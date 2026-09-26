"""Librería de estilo Kin Analytics para presentaciones (python-pptx).

Estilo tomado del deck de referencia de Kin ("Customer Potential Model", 2026):
- 16:9 de 10 × 5.625 in. Portadas y separadores en negro con título blanco grande, subtítulo en lima y un corte
  lima en la esquina superior derecha. Láminas de contenido en blanco.
- Etiqueta (kicker) gris en mayúsculas arriba del título; título negro en negrita; subtítulo gris.
- Tarjetas gris claro (F4F4F4) o negras; números grandes para cifras clave; tablas con encabezado negro.
- Mensaje clave en barra negra con texto lima (o lima con texto negro). Acento único: lima E5FF01.
- Pie: "Kin Analytics®" · "Let knowledge in." · "KinAnalytics.com" sobre franja gris clara.
Fuentes de marca: Albert Sans (texto) y Space Grotesk (títulos). No están instaladas en todas las máquinas y
PowerPoint las sustituye; por defecto se usa Arial (FUENTE_*). Cambiar aquí si el equipo las tiene.
"""
from pptx import Presentation
from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ── paleta y tipografía (única fuente de verdad del estilo) ──
NEGRO, NEGRO_2, GRIS, GRIS_MEDIO = "0A0A0A", "222326", "666666", "999EA6"
GRIS_CLARO, FONDO_CARD, BLANCO, LIMA, VERDE = "E0E0E0", "F4F4F4", "FFFFFF", "E5FF01", "22C55E"
FUENTE_TITULO = "Arial"          # marca: "Space Grotesk"
FUENTE_TEXTO = "Arial"           # marca: "Albert Sans"
ANCHO, ALTO = 10.0, 5.625
MARGEN = 0.45


def rgb(h):
    return RGBColor.from_string(h)


def nueva():
    """Presentación vacía 16:9 (10 × 5.625 in)."""
    p = Presentation()
    p.slide_width, p.slide_height = Inches(ANCHO), Inches(ALTO)
    return p


def _vacia(prs, fondo=BLANCO):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(fondo)
    return s


def texto(slide, x, y, w, h, contenido, tam=12, color=NEGRO, negrita=False, fuente=None, alinear=PP_ALIGN.LEFT,
          vertical=MSO_ANCHOR.TOP, interlineado=1.05, cursiva=False):
    """Caja de texto sin márgenes internos. `contenido` = str o lista de str (un párrafo por elemento) o de
    tuplas (str, dict de formato) para mezclar estilos en líneas distintas."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = vertical
    lineas = contenido if isinstance(contenido, list) else str(contenido).split("\n")   # salto de línea = párrafo nuevo
    for i, linea in enumerate(lineas):
        txt, fmt = (linea if isinstance(linea, tuple) else (linea, {}))
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = fmt.get("alinear", alinear)
        par.line_spacing = interlineado
        par.space_after = Pt(fmt.get("espacio", 0))
        r = par.add_run()
        r.text = str(txt)
        f = r.font
        f.name = fmt.get("fuente", fuente or FUENTE_TEXTO)
        f.size = Pt(fmt.get("tam", tam))
        f.bold = fmt.get("negrita", negrita)
        f.italic = fmt.get("cursiva", cursiva)
        f.color.rgb = rgb(fmt.get("color", color))
    return tb


def caja(slide, x, y, w, h, relleno=FONDO_CARD, borde=None, redondeo=0.06, forma=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(relleno)
    if borde:
        sh.line.color.rgb = rgb(borde)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = redondeo
    return sh


def _corte_lima(slide):
    """Corte lima en la esquina superior derecha (motivo de portadas y separadores)."""
    ff = slide.shapes.build_freeform(Inches(ANCHO - 0.95), Inches(0))
    ff.add_line_segments([(Inches(ANCHO), Inches(0)), (Inches(ANCHO), Inches(1.55)), (Inches(ANCHO - 0.35), Inches(0.72))])
    sh = ff.convert_to_shape()
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(LIMA)
    sh.line.fill.background()


def _pie(slide, oscuro=False):
    y = ALTO - 0.32
    if not oscuro:
        caja(slide, 0, y, ANCHO, 0.32, relleno=FONDO_CARD, forma=MSO_SHAPE.RECTANGLE)
    col, col2 = (BLANCO, GRIS_MEDIO) if oscuro else (NEGRO, GRIS_MEDIO)
    texto(slide, MARGEN, y + 0.09, 1.3, 0.16, "Kin Analytics®", tam=8.5, color=col, negrita=True, fuente=FUENTE_TITULO)
    texto(slide, MARGEN + 1.25, y + 0.1, 2.0, 0.14, "Let knowledge in.", tam=6.5, color=col2)
    texto(slide, ANCHO - MARGEN - 1.6, y + 0.1, 1.6, 0.14, "KinAnalytics.com", tam=6.5, color=col2, alinear=PP_ALIGN.RIGHT)


def portada(prs, titulo, subtitulo, bajada="", fecha_cliente="", confidencial=True):
    s = _vacia(prs, NEGRO)
    _corte_lima(s)
    if confidencial:
        texto(s, MARGEN, 0.3, 3, 0.4, ["Kin Analytics ©2026", "CONFIDENCIAL"], tam=6.5, color=GRIS_MEDIO)
    texto(s, MARGEN, 1.0, 8.2, 1.9, titulo, tam=44, color=BLANCO, negrita=True, fuente=FUENTE_TITULO, interlineado=0.92,
          vertical=MSO_ANCHOR.BOTTOM)
    texto(s, MARGEN, 3.0, 8.5, 0.35, subtitulo, tam=15, color=LIMA, fuente=FUENTE_TITULO)
    if bajada:
        texto(s, MARGEN, 3.55, 7.5, 0.5, bajada, tam=9.5, color=GRIS_MEDIO)
    texto(s, MARGEN, ALTO - 0.55, 5, 0.2, fecha_cliente, tam=8.5, color=LIMA)
    texto(s, ANCHO - 3.0, ALTO - 0.72, 2.6, 0.45, "Kin Analytics®", tam=22, color=BLANCO, fuente=FUENTE_TITULO, alinear=PP_ALIGN.RIGHT)
    return s


def separador(prs, linea1, linea2=""):
    s = _vacia(prs, NEGRO)
    _corte_lima(s)
    texto(s, MARGEN, 0.3, 3, 0.4, ["Kin Analytics", "CONFIDENCIAL"], tam=6.5, color=GRIS_MEDIO)
    texto(s, MARGEN, 1.7, 9, 0.8, linea1, tam=38, color=BLANCO, negrita=True, fuente=FUENTE_TITULO)
    if linea2:
        texto(s, MARGEN, 2.45, 9, 0.8, linea2, tam=38, color=LIMA, negrita=True, fuente=FUENTE_TITULO)
    _pie(s, oscuro=True)
    return s


def lamina(prs, titulo, kicker="", subtitulo="", notas=""):
    """Lámina de contenido en blanco. Devuelve (slide, y_inicio_contenido)."""
    s = _vacia(prs, BLANCO)
    y = 0.28
    if kicker:
        texto(s, MARGEN, y, 9, 0.16, kicker.upper(), tam=7, color=GRIS, negrita=True)
        y += 0.2
    texto(s, MARGEN, y, 9.1, 0.5, titulo, tam=21, negrita=True, fuente=FUENTE_TITULO, interlineado=0.95)
    y += 0.5 if len(titulo) <= 56 else 0.85          # ~56 caracteres por línea a 21 pt en 9.1 in
    if subtitulo:
        texto(s, MARGEN, y, 9.1, 0.35, subtitulo, tam=9.5, color=GRIS)
        y += 0.38
    _pie(s)
    if notas:
        s.notes_slide.notes_text_frame.text = notas
    return s, y + 0.1


def tarjeta(slide, x, y, w, h, titulo, cuerpo="", numero=None, oscura=False, acento=False, compacta=False):
    """Tarjeta con número opcional en círculo lima, título en negrita y texto (compacta = texto más pegado al título)."""
    fondo = NEGRO if oscura else (LIMA if acento else FONDO_CARD)
    caja(slide, x, y, w, h, relleno=fondo)
    c_tit = BLANCO if oscura else NEGRO
    c_txt = GRIS_CLARO if oscura else (NEGRO if acento else GRIS)
    x0 = x + 0.14
    if numero is not None:
        circ = caja(slide, x0, y + 0.14, 0.3, 0.3, relleno=LIMA if not acento else NEGRO, forma=MSO_SHAPE.OVAL)
        texto(slide, x0, y + 0.19, 0.3, 0.2, str(numero), tam=9, negrita=True, color=NEGRO if not acento else LIMA,
              alinear=PP_ALIGN.CENTER)
        x0 += 0.4
    texto(slide, x0, y + 0.16, w - (x0 - x) - 0.12, 0.3, titulo, tam=10.5, negrita=True, color=c_tit, fuente=FUENTE_TITULO)
    if cuerpo:
        dy = 0.42 if compacta else 0.55
        texto(slide, x + 0.14, y + dy, w - 0.28, h - dy - 0.06, cuerpo, tam=8.5, color=c_txt, interlineado=1.1)


def cifra(slide, x, y, w, valor, etiqueta, detalle="", h=1.0, oscura=False):
    """Cifra grande (estilo 'stat') dentro de tarjeta."""
    caja(slide, x, y, w, h, relleno=NEGRO if oscura else FONDO_CARD)
    texto(slide, x + 0.15, y + 0.1, w - 0.3, 0.5, valor, tam=24, negrita=True, fuente=FUENTE_TITULO,
          color=LIMA if oscura else NEGRO)
    texto(slide, x + 0.15, y + 0.58, w - 0.3, 0.2, etiqueta, tam=8.5, negrita=True, color=BLANCO if oscura else NEGRO)
    if detalle:
        texto(slide, x + 0.15, y + 0.77, w - 0.3, max(h - 0.8, 0.18), detalle, tam=7, color=GRIS_MEDIO if oscura else GRIS)


def mensaje(slide, contenido, y=None, invertido=False):
    """Barra de mensaje clave (negra con texto lima; invertido = lima con texto negro)."""
    y = ALTO - 0.32 - 0.5 if y is None else y
    caja(slide, MARGEN, y, ANCHO - 2 * MARGEN, 0.38, relleno=LIMA if invertido else NEGRO, redondeo=0.12)
    texto(slide, MARGEN + 0.18, y + 0.1, ANCHO - 2 * MARGEN - 0.36, 0.2, contenido, tam=9, negrita=True,
          color=NEGRO if invertido else LIMA)


def pd_es_numerica(serie) -> bool:
    """True si la columna es numérica (se alinea a la derecha; el texto va a la izquierda)."""
    return all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in serie)


def tabla(slide, df, x, y, w, anchos=None, alto_fila=0.26, tam=7.5, resaltar=None, formato=None):
    """Tabla con encabezado negro. `resaltar` = índices de fila (0 = primera de datos) con fondo lima.
    `formato` = {columna: función(valor) -> str}."""
    filas, cols = len(df) + 1, len(df.columns)
    t = slide.shapes.add_table(filas, cols, Inches(x), Inches(y), Inches(w), Inches(alto_fila * filas)).table
    anchos = anchos or [w / cols] * cols
    for j, a in enumerate(anchos):
        t.columns[j].width = Inches(a)
    formato = formato or {}
    for i in range(filas):
        t.rows[i].height = Inches(alto_fila)
        for j, col in enumerate(df.columns):
            celda = t.cell(i, j)
            celda.margin_left = celda.margin_right = Inches(0.06)
            celda.margin_top = celda.margin_bottom = Inches(0.02)
            if i == 0:
                valor, fondo, color, negr = str(col), NEGRO, BLANCO, True
            else:
                v = df.iloc[i - 1, j]
                valor = formato.get(col, lambda z: f"{z:,.0f}" if isinstance(z, (int, float)) and not isinstance(z, bool) else str(z))(v)
                fila_res = resaltar is not None and (i - 1) in resaltar
                fondo = LIMA if fila_res else (BLANCO if i % 2 else FONDO_CARD)
                color, negr = NEGRO, fila_res or j == 0
            celda.fill.solid()
            celda.fill.fore_color.rgb = rgb(fondo)
            tf = celda.text_frame
            tf.text = valor
            p = tf.paragraphs[0]
            numerica = i > 0 and isinstance(df.iloc[i - 1, j], (int, float)) and not isinstance(df.iloc[i - 1, j], bool)
            if i == 0:
                numerica = pd_es_numerica(df.iloc[:, j])
            p.alignment = PP_ALIGN.RIGHT if (j > 0 and numerica) else PP_ALIGN.LEFT
            f = p.runs[0].font
            f.size, f.bold, f.name = Pt(tam), negr, FUENTE_TEXTO
            f.color.rgb = rgb(color)
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
    return t


def barras(slide, x, y, w, h, categorias, series: dict, titulo="", horizontal=True, apiladas=False, formato="0",
           colores=None, eje_max=None, leyenda=None):
    """Gráfico de barras nativo (editable en PowerPoint) con etiquetas de datos y ejes discretos."""
    datos = CategoryChartData()
    datos.categories = list(categorias)
    for nombre, valores in series.items():
        datos.add_series(nombre, [float(v) for v in valores])
    if horizontal:
        tipo = XL_CHART_TYPE.BAR_STACKED if apiladas else XL_CHART_TYPE.BAR_CLUSTERED
    else:
        tipo = XL_CHART_TYPE.COLUMN_STACKED if apiladas else XL_CHART_TYPE.COLUMN_CLUSTERED
    graf = slide.shapes.add_chart(tipo, Inches(x), Inches(y), Inches(w), Inches(h), datos).chart
    colores = colores or [NEGRO, LIMA, GRIS_MEDIO, VERDE, GRIS_CLARO]
    graf.has_title = bool(titulo)
    if titulo:
        graf.chart_title.text_frame.text = titulo
        r = graf.chart_title.text_frame.paragraphs[0].runs[0].font
        r.size, r.bold, r.name = Pt(9), True, FUENTE_TITULO
        r.color.rgb = rgb(NEGRO)
    graf.has_legend = len(series) > 1 if leyenda is None else leyenda
    if graf.has_legend:
        graf.legend.include_in_layout = False
        graf.legend.font.size = Pt(7)
        graf.legend.font.color.rgb = rgb(GRIS)
    plot = graf.plots[0]
    plot.gap_width = 60
    if apiladas:
        plot.overlap = 100
    for i, serie in enumerate(plot.series):
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = rgb(colores[i % len(colores)])
        serie.data_labels.show_value = True
        serie.data_labels.number_format = formato
        serie.data_labels.number_format_is_linked = False
        serie.data_labels.font.size = Pt(7)
        serie.data_labels.font.color.rgb = rgb(BLANCO if colores[i % len(colores)] in (NEGRO, NEGRO_2) and apiladas else NEGRO)
        serie.data_labels.position = XL_LABEL_POSITION.CENTER if apiladas else XL_LABEL_POSITION.OUTSIDE_END
        if apiladas:                         # segmentos < 3% de la barra más larga sin etiqueta: un "0" sobre una franja invisible no se lee
            umbral = 0.03 * max(sum(abs(float(vs[k])) for vs in series.values()) for k in range(len(categorias)))
            for j, v in enumerate(list(series.values())[i]):
                if abs(float(v)) < umbral:
                    dl = serie.points[j].data_label._get_or_add_dLbl()
                    for hijo in list(dl):
                        if hijo.tag != qn("c:idx"):
                            dl.remove(hijo)
                    etree.SubElement(dl, qn("c:delete")).set("val", "1")
    for eje in (graf.category_axis, graf.value_axis):
        eje.tick_labels.font.size = Pt(7)
        eje.tick_labels.font.color.rgb = rgb(GRIS)
        eje.format.line.color.rgb = rgb(GRIS_CLARO)
    graf.value_axis.has_major_gridlines = False
    graf.value_axis.visible = False
    if eje_max is not None:
        graf.value_axis.maximum_scale = eje_max
        graf.value_axis.minimum_scale = 0
    if horizontal:
        graf.category_axis.reverse_order = True
    return graf


def cierre(prs, frase="Let knowledge in.", contacto=""):
    s = _vacia(prs, NEGRO)
    _corte_lima(s)
    texto(s, MARGEN, 1.9, 9, 0.8, frase, tam=40, color=BLANCO, negrita=True, fuente=FUENTE_TITULO)
    texto(s, MARGEN, 2.8, 9, 0.3, "KinAnalytics.com", tam=14, color=LIMA, fuente=FUENTE_TITULO)
    if contacto:
        texto(s, MARGEN, 3.3, 9, 0.3, contacto, tam=9, color=GRIS_MEDIO)
    texto(s, ANCHO - 3.0, ALTO - 0.72, 2.6, 0.45, "Kin Analytics®", tam=22, color=BLANCO, fuente=FUENTE_TITULO, alinear=PP_ALIGN.RIGHT)
    return s


def notas(slide, texto_notas):
    """Notas para el orador (guion) de una lámina."""
    slide.notes_slide.notes_text_frame.text = texto_notas
