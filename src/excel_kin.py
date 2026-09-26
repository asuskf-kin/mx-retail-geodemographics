"""Formato corporativo Kin para entregables Excel (openpyxl).

Misma identidad que `deck_kin`: negro 0A0A0A, acento lima E5FF01, grises, Arial. Funciones:
- `portada(wb, titulo, subtitulo, filas_info, indice)`: hoja inicial con título, datos del estudio e índice de hojas.
- `hoja_tabla(wb, nombre, df, ...)`: tabla con encabezado negro, filtros, paneles fijos, anchos, formatos numéricos,
  columnas de semáforo / clúster coloreadas y barras de datos opcionales.
- `hoja_texto(wb, nombre, bloques)`: hojas de lectura (guion, accionables, diccionario) con títulos y párrafos.
- `hoja_tabla_2niveles(wb, nombre, df, ...)`: tabla con encabezado de 2 niveles (grupo | subcolumna, p. ej. categoría |
  HHs, %, Indice), bloques de color por grupo, índice resaltado (≥ 105 lima) y filas alternadas.
- `libro()`: Workbook vacío con Arial 9 como fuente por defecto (evita dar formato celda por celda en tablas grandes).
"""
import pandas as pd
import openpyxl
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NEGRO, GRIS, GRIS_MEDIO, GRIS_CLARO, FONDO, LIMA, BLANCO = "0A0A0A", "666666", "999EA6", "E0E0E0", "F4F4F4", "E5FF01", "FFFFFF"
SEMAFORO = {"verde": "22C55E", "amarillo": "EDA100", "rojo": "E34948"}
CLUSTER = {"HH": NEGRO, "HL": GRIS, "LH": "22C55E", "LL": GRIS_MEDIO}
F = "Arial"
_linea = Side(style="thin", color=GRIS_CLARO)


def _relleno(c):
    return PatternFill("solid", fgColor=c)


def libro():
    """Workbook vacío con la fuente Kin (Arial 9) como estilo Normal."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    normal = next(st for st in wb._named_styles if st.name == "Normal")
    normal.font = Font(name=F, size=9, color=NEGRO)
    return wb


def portada(wb, titulo, subtitulo, filas_info, indice):
    """Primera hoja: título, subtítulo, bloque de datos del estudio (lista de (etiqueta, valor)) e índice (lista de (hoja, qué contiene))."""
    ws = wb.create_sheet("Portada", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 95
    for r in range(1, 7):
        for c in range(1, 4):
            ws.cell(r, c).fill = _relleno(NEGRO)
    ws["B2"] = "Kin Analytics®"
    ws["B2"].font = Font(name=F, size=10, bold=True, color=LIMA)
    ws["B3"] = titulo
    ws["B3"].font = Font(name=F, size=24, bold=True, color=BLANCO)
    ws["B5"] = subtitulo
    ws["B5"].font = Font(name=F, size=12, color=LIMA)
    r = 8
    ws.cell(r, 2, "DATOS DEL ESTUDIO").font = Font(name=F, size=9, bold=True, color=GRIS)
    for etq, val in filas_info:
        r += 1
        ws.cell(r, 2, etq).font = Font(name=F, size=10, bold=True, color=NEGRO)
        ws.cell(r, 3, val).font = Font(name=F, size=10, color=NEGRO)
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    r += 2
    ws.cell(r, 2, "CONTENIDO").font = Font(name=F, size=9, bold=True, color=GRIS)
    for hoja, que in indice:
        r += 1
        c = ws.cell(r, 2, hoja)
        c.hyperlink = f"#'{hoja}'!A1"
        c.font = Font(name=F, size=10, bold=True, color=NEGRO, underline="single")
        ws.cell(r, 3, que).font = Font(name=F, size=10, color=GRIS)
    r += 2
    ws.cell(r, 2, "Let knowledge in. · KinAnalytics.com · CONFIDENCIAL").font = Font(name=F, size=8, color=GRIS_MEDIO)
    return ws


def hoja_tabla(wb, nombre, df, formatos=None, anchos=None, semaforo_col=None, cluster_col=None, barras=None,
               titulo=None, nota=None, indice=False, ajustar=False, escala=None):
    """Tabla con estilo Kin. formatos = {columna: formato Excel}; barras = columnas con barra de datos lima;
    ajustar = texto largo con salto de línea; escala = columnas con escala divergente (gris ← 0 → lima)."""
    ws = wb.create_sheet(nombre)
    ws.sheet_view.showGridLines = False
    fila0 = 1
    if titulo:
        ws.cell(1, 1, titulo).font = Font(name=F, size=13, bold=True, color=NEGRO)
        if nota:
            ws.cell(2, 1, nota).font = Font(name=F, size=9, italic=True, color=GRIS)
        fila0 = 4 if nota else 3
    d = df.reset_index() if indice else df
    cols = list(d.columns)
    for j, col in enumerate(cols, 1):
        c = ws.cell(fila0, j, str(col))
        c.font = Font(name=F, size=9, bold=True, color=BLANCO)
        c.fill = _relleno(NEGRO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    formatos = formatos or {}
    for i, fila in enumerate(d.itertuples(index=False), fila0 + 1):
        for j, v in enumerate(fila, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(i, j, v)
            c.font = Font(name=F, size=9, color=NEGRO)
            c.border = Border(bottom=_linea)
            if cols[j - 1] in formatos:
                c.number_format = formatos[cols[j - 1]]
            if ajustar:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    ultima = fila0 + len(d)
    ws.row_dimensions[fila0].height = 30
    for j, col in enumerate(cols, 1):
        letra = get_column_letter(j)
        ancho = (anchos or {}).get(col) or min(max(len(str(col)), *(len(str(x)) for x in d[col].head(200))) + 2, 45)
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = ws.cell(fila0 + 1, 2)
    if len(d):
        ws.auto_filter.ref = f"A{fila0}:{get_column_letter(len(cols))}{ultima}"
    rango = lambda col: f"{get_column_letter(cols.index(col) + 1)}{fila0 + 1}:{get_column_letter(cols.index(col) + 1)}{ultima}"
    if semaforo_col and semaforo_col in cols:
        for valor, color in SEMAFORO.items():
            ws.conditional_formatting.add(rango(semaforo_col), CellIsRule(operator="equal", formula=[f'"{valor}"'],
                                          fill=_relleno(color), font=Font(name=F, bold=True, color=BLANCO)))
    if cluster_col and cluster_col in cols:
        for valor, color in CLUSTER.items():
            ws.conditional_formatting.add(rango(cluster_col), CellIsRule(operator="equal", formula=[f'"{valor}"'], fill=_relleno(color),
                                          font=Font(name=F, bold=True, color=LIMA if valor == "HH" else BLANCO)))
    for col in (barras or []):
        if col in cols:
            ws.conditional_formatting.add(rango(col), DataBarRule(start_type="min", end_type="max", color="C8D400"))
    for col in (escala or []):
        if col in cols:
            ws.conditional_formatting.add(rango(col), ColorScaleRule(start_type="num", start_value=-0.3, start_color="999EA6",
                                                                     mid_type="num", mid_value=0, mid_color=BLANCO,
                                                                     end_type="num", end_value=0.3, end_color=LIMA))
    if ajustar:                                   # alto de fila según el texto más largo respecto al ancho de su columna
        for i in range(fila0 + 1, ultima + 1):
            lineas = max(len(str(ws.cell(i, j).value or "")) / max(ws.column_dimensions[get_column_letter(j)].width - 1, 1)
                         for j in range(1, len(cols) + 1))
            ws.row_dimensions[i].height = max(15, 12.5 * (int(lineas) + 1))
    return ws


def hoja_tabla_2niveles(wb, nombre, df, titulo=None, nota=None, formatos=None, anchos=None, colores_grupo=None,
                        fijar_cols=2, indice_alto=105, indice_bajo=95, nombre_indice="Indice"):
    """Tabla con encabezado de 2 niveles (columnas MultiIndex: grupo, subcolumna).
    formatos = {(grupo, sub) o sub: formato}; colores_grupo = {grupo: color del bloque}; las subcolumnas `nombre_indice`
    se resaltan en lima si ≥ indice_alto y en gris si ≤ indice_bajo (mismo criterio que el deck)."""
    ws = wb.create_sheet(nombre)
    ws.sheet_view.showGridLines = False
    fila0 = 1
    if titulo:
        ws.cell(1, 1, titulo).font = Font(name=F, size=13, bold=True, color=NEGRO)
        if nota:
            ws.cell(2, 1, nota).font = Font(name=F, size=9, italic=True, color=GRIS)
        fila0 = 4
    cols = list(df.columns)
    formatos, colores_grupo = formatos or {}, colores_grupo or {}
    blanco = Side(style="thin", color=BLANCO)
    # nivel 1: grupos (celdas combinadas) · nivel 2: subcolumnas
    j = 1
    while j <= len(cols):
        g = cols[j - 1][0]
        k = j
        while k < len(cols) and cols[k][0] == g:
            k += 1
        c = ws.cell(fila0, j, g)
        c.font = Font(name=F, size=9, bold=True, color=LIMA if colores_grupo.get(g, NEGRO) == NEGRO else BLANCO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for jj in range(j, k + 1):
            ws.cell(fila0, jj).fill = _relleno(colores_grupo.get(g, NEGRO))
            ws.cell(fila0, jj).border = Border(left=blanco if jj == j else None)
        if k > j:
            ws.merge_cells(start_row=fila0, start_column=j, end_row=fila0, end_column=k)
        j = k + 1
    for j, (g, sub) in enumerate(cols, 1):
        c = ws.cell(fila0 + 1, j, sub)
        c.font = Font(name=F, size=9, bold=True, color=NEGRO)
        c.fill = _relleno(GRIS_CLARO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila0].height = 30
    ws.row_dimensions[fila0 + 1].height = 42
    fmt = [formatos.get(col, formatos.get(col[1])) for col in cols]
    for i, fila in enumerate(df.itertuples(index=False), fila0 + 2):
        for j, v in enumerate(fila, 1):
            if isinstance(v, float) and pd.isna(v):
                continue
            c = ws.cell(i, j, v.item() if hasattr(v, "item") else v)
            if fmt[j - 1]:
                c.number_format = fmt[j - 1]
    ultima = fila0 + 1 + len(df)
    for j, (g, sub) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = (anchos or {}).get(sub) or (anchos or {}).get((g, sub)) or             min(max(len(str(sub)) * 0.9, *(len(str(x)) for x in df[(g, sub)].head(200))) + 2, 40)
    ws.freeze_panes = ws.cell(fila0 + 2, fijar_cols + 1)
    fin = get_column_letter(len(cols))
    ws.auto_filter.ref = f"A{fila0 + 1}:{fin}{ultima}"
    for j, (g, sub) in enumerate(cols, 1):
        if sub == nombre_indice:
            r = f"{get_column_letter(j)}{fila0 + 2}:{get_column_letter(j)}{ultima}"
            ws.conditional_formatting.add(r, CellIsRule(operator="greaterThanOrEqual", formula=[str(indice_alto)], fill=_relleno(LIMA),
                                                        font=Font(name=F, bold=True, color=NEGRO)))
            ws.conditional_formatting.add(r, CellIsRule(operator="lessThanOrEqual", formula=[str(indice_bajo)], font=Font(name=F, color=GRIS_MEDIO)))
    ws.conditional_formatting.add(f"A{fila0 + 2}:{fin}{ultima}", FormulaRule(formula=[f"MOD(ROW(),2)=0"], fill=_relleno(FONDO)))
    return ws


def hoja_texto(wb, nombre, bloques, ancho=110):
    """Hoja de lectura: bloques = lista de (titulo, [párrafos]) ."""
    ws = wb.create_sheet(nombre)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = ancho
    r = 1
    for titulo, parrafos in bloques:
        r += 1
        c = ws.cell(r, 2, titulo)
        c.font = Font(name=F, size=12, bold=True, color=NEGRO)
        c.fill = _relleno(LIMA)
        for p in parrafos:
            r += 1
            c = ws.cell(r, 2, p)
            c.font = Font(name=F, size=10, color=NEGRO)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = max(15, 15 * (len(str(p)) // ancho + 1))
        r += 1
    return ws


def estilo_formulas(ws, rango_celdas, formato=None, negrita=False, relleno=None):
    """Aplica formato a celdas con fórmulas escritas a mano (hoja Resumen)."""
    for fila in ws[rango_celdas]:
        for c in fila:
            c.font = Font(name=F, size=10, bold=negrita, color=NEGRO)
            if formato:
                c.number_format = formato
            if relleno:
                c.fill = _relleno(relleno)
            c.border = Border(bottom=_linea)
