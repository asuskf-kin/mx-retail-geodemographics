"""Versión KO del 06 (sin puntos): mapa y Excel reducido por hexágono H3 en res 9 (clúster NSE), res 8 y res 7.

Uso (desde la raíz del repo; lo corre solo src/correr.py al terminar bien el 06):
    uv run python .claude/skills/golden-stores-kin/scripts/version_ko.py <ciudad>
    uv run python .claude/skills/golden-stores-kin/scripts/version_ko.py <carpeta cliente> [--estados 09,15]
        (p. ej. outputs/guadalajara/arca/cliente: busca ahí el mapa 06_mapa_qa_letras_*.html y el *reducido*.xlsx NORMALES y escribe
        al lado sus _KO; usuario 2026-10-23: "te voy a mandar esta ruta para que actualices mapa y excel")
    uv run python .claude/skills/golden-stores-kin/scripts/version_ko.py --archivos <mapa.html> <reducido.xlsx> <mapa_KO.html> <reducido_KO.xlsx>
        [--estados 09,15]   (entregables sueltos con el mismo formato del 06, p. ej. FEMSA CDMX descargado; --estados = claves INEGI
        de las entidades: si el reducido o el mapa no traen código postal, se calcula con el DENUE de esas entidades)

Lee los entregables del 06 de la ciudad (no los cambia) y escribe al lado, con sufijo _KO:
- outputs/<ciudad>/<cliente>/cliente/06_mapa_qa_letras_<cliente>_zm_<ciudad>_KO.html: el mismo mapa (con la plantilla vigente de src/mapas.py)
  sin capas de puntos (PDV, Rappi y sus radios), sin los datos de esos puntos y sin filtro por subcanal. Cada hexágono muestra la
  letra final y la prioridad que PREDOMINAN entre sus PDV (empate → la mejor: HH, HL, LH, LL · P1 … P4), con el conteo por letra en
  el tooltip, en res 9 (= clúster NSE), res 8 y res 7: chips "Hexágono" para elegir la resolución y búsqueda por id del hexágono.
  Los pasos que solo tenían puntos (Letra 2, letras base) se quitan.
- outputs/<ciudad>/<cliente>/cliente/06_letras_nse_ventas_tradicional_reducido_<cliente>_zm_<ciudad>_KO.xlsx: misma plantilla del
  reducido (bloques de color, formatos, 3.ª fila con la fuente, Diccionario), una fila por PDV, columnas
  Hexágono res 9 · res 8 · res 7 · Código postal (zipcode, si el reducido lo trae) · pos_id · Subcanal · resto igual (sin Nombre, Latitud ni
  Longitud ni la columna repetida del clúster) y al final, en Clasificación, Letras · Acción · Prioridad del hexágono por resolución.
Res 9 es el clúster NSE del PDV (hexágono que lo contiene); res 8 y res 7 son los hexágonos que contienen al PDV (h3.latlng_to_cell).
"""
import json
import os
import sys
from collections import Counter
from pathlib import Path

import h3
import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parents[4]
_ARG = sys.argv[1] if len(sys.argv) > 1 else ""
CARPETA = Path(_ARG).resolve() if _ARG and _ARG != "--archivos" and Path(_ARG).is_dir() else None   # .../outputs/<ciudad>/<bottler>/cliente
CIUDAD = (CARPETA.parts[-3] if CARPETA else _ARG) if _ARG and _ARG != "--archivos" else os.environ.get("CIUDAD", "merida")
os.environ["CIUDAD"] = CIUDAD
sys.path.insert(0, str(RAIZ / "src"))
try:
    import config as C      # noqa: E402  (lee CIUDAD al cargarse)
    CONFIGURADA = True
except KeyError:            # ciudad que no está en config.CIUDADES (p. ej. una carpeta de FEMSA CDMX): se carga con la de defecto
    os.environ["CIUDAD"] = "merida"
    import config as C      # noqa: E402
    CONFIGURADA = False
import excel_kin as X       # noqa: E402
import mapas as MP          # noqa: E402
import codigos_postales as CPS   # noqa: E402

DENUE_EXTRA = None          # --estados: DENUE para poner el código postal a entregables sueltos que no lo traen

RES = [9, 8, 7]                                              # 9 = clúster NSE
AREA = {9: "≈ 0.1 km²", 8: "≈ 0.7 km²", 7: "≈ 5.2 km²"}     # área media del hexágono H3
LET, PRI = ["HH", "HL", "LH", "LL"], ["P1", "P2", "P3", "P4"]
ACCION = {"HH": "Atacar", "HL": "Bloquear", "LH": "Fortalecer", "LL": "Mantener"}
SIN = "sin NSE (sin hogares a 300 m)"


def moda(vals, orden):
    """La que más se repite; empate → la primera de `orden` (la mejor)."""
    c = Counter(v for v in vals if v in orden)
    return min(c, key=lambda k: (-c[k], orden.index(k))) if c else None


def hex_res(cl9, lat, lon, r):
    """Hexágono del PDV en la resolución r: res 9 = su clúster NSE; res 8 y 7 = el que contiene al punto."""
    if r == 9:
        return cl9 if isinstance(cl9, str) and cl9 else None
    try:
        return h3.latlng_to_cell(float(lat), float(lon), r)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------- mapa
def _media(vals, pesos):
    """Media ponderada sin faltantes (None si no hay datos)."""
    par = [(v, w) for v, w in zip(vals, pesos) if v is not None and w]
    s = sum(w for _, w in par)
    return round(sum(v * w for v, w in par) / s, 2) if s else None


def tablas_res(D) -> dict:
    """Una tabla por resolución con TODOS los datos del hexágono, para que cada paso del mapa se vea en la resolución elegida.
    Hogares y NSE de los hogares (y Censo): de los hexágonos res 10 del 06 (suma y media ponderada por hogares) · clúster NSE
    (nivel INEGI de su buffer, AltScore, nivel final): res 9 = el clúster; res 8/7 = media de sus clústeres ponderada por PDV ·
    letras y prioridad: las que predominan entre sus PDV · share Whisp: el del hexágono res 6 que lo contiene."""
    H10, CL, WH = D["hex"]["hex"], D["hex"]["cl"], D["hex"].get("wh")
    i10 = {k: H10["campos"].index(k) for k in H10["campos"]}
    icl = {k: CL["campos"].index(k) for k in CL["campos"]}
    P = D["datasets"]["pdv"]
    ip = {k: P["campos"].index(k) for k in ("cl", "lat", "lon", "letras", "prioridad")}
    izip = P["campos"].index("zip") if "zip" in P["campos"] else None      # código postal del PDV (06 desde 2026-10-22)
    tablas = {}
    for r in RES:
        hog, clu, pdv = {}, {}, {}
        for h, v in H10["filas"].items():
            hog.setdefault(h3.cell_to_parent(h, r), []).append(v)
        for h, v in CL["filas"].items():
            clu.setdefault(h if r == 9 else h3.cell_to_parent(h, r), []).append(v)
        for f in P["filas"]:
            h = hex_res(f[ip["cl"]], f[ip["lat"]], f[ip["lon"]], r) if f[ip["cl"]] else None
            if h:
                pdv.setdefault(h, []).append(f)
        filas = {}
        for h in set(hog) | set(clu) | set(pdv):
            a, c_, p_ = hog.get(h, []), clu.get(h, []), pdv.get(h, [])
            w = [v[i10["hog"]] or 0 for v in a]
            wc = [v[icl["n_pdv"]] or 0 for v in c_]
            cl = Counter(f[ip["letras"]] for f in p_ if f[ip["letras"]] in LET)
            cp = Counter(f[ip["prioridad"]] for f in p_ if f[ip["prioridad"]] in PRI)
            n = sum(cl.values())
            letra = moda([f[ip["letras"]] for f in p_], LET)
            share, clase = None, None
            if WH:
                sw = WH["filas"].get(h3.cell_to_parent(h, 6))
                if sw:
                    share, clase = sw
            zc = Counter(f[izip] for f in p_ if izip is not None and isinstance(f[izip], str) and f[izip].isdigit())
            filas[h] = {
                "zip": (min(zc, key=lambda k: (-zc[k], k)) if zc else None), "n_zip": len(zc) or None,
                "zips": " · ".join(f"{k} ({v})" for k, v in sorted(zc.items(), key=lambda kv: (-kv[1], kv[0]))[:6]) or None,
                "hog_h": round(sum(w), 0) if a else None, "N_h": _media([v[i10["N"]] for v in a], w),
                "p_bajo": _media([v[i10["p_bajo"]] for v in a], w), "p_medio": _media([v[i10["p_medio"]] for v in a], w),
                "p_alto": _media([v[i10["p_alto"]] for v in a], w), "censo_h": _media([v[i10["censo"]] for v in a], w),
                "n_cl": len(c_) or None, "N_cl": _media([v[icl["N"]] for v in c_], wc), "alt": _media([v[icl["alt"]] for v in c_], wc),
                "Nf": _media([v[icl["Nf"]] for v in c_], wc),
                "n_pdv": len(p_), "n_letra": n, **{f"n_{k}": cl[k] for k in LET},
                **{f"p_{k}": (round(100 * cl[k] / n, 1) if n else None) for k in LET},
                "letras": letra, "accion": ACCION.get(letra), "prioridad": moda([f[ip["prioridad"]] for f in p_], PRI),
                **{f"n_{k}": cp[k] for k in PRI}, "share": share, "share_clase": clase}
        campos = list(next(iter(filas.values())))
        tablas[r] = {"campos": campos, "filas": {h: [v[k] for k in campos] for h, v in filas.items()}}
    return tablas


PALETA_ZIP = ["#2a78d6", "#eda100", "#1baf7a", "#e34948", "#7b5ea7", "#0a0a0a", "#f08bb5", "#8c6d31", "#4fb3d9", "#a3c639", "#c65a2e", "#5c6bc0"]


def ESC_ZIP(tablas):
    """Color por código postal: códigos consecutivos (vecinos) con colores distintos; la leyenda solo explica cómo leerlo."""
    zips = sorted({v[0] for t in tablas.values() for v in t["filas"].values() if v[0]})
    return {"tipo": "categorica", "colores": {z: PALETA_ZIP[int(z) % len(PALETA_ZIP)] for z in zips},
            "leyenda": f"Color = código postal ({len(zips):,} códigos). El número está en el tooltip y en la ficha; búscalo arriba."}


def mapa_ko(src: Path, dst: Path):
    t = src.read_text(encoding="utf-8")
    a = t.index("const D = ") + len("const D = ")
    b = t.index(";\ndocument.getElementById('t')", a)
    D = json.loads(t[a:b])
    P = D["datasets"]["pdv"]
    if "zip" not in P["campos"] and DENUE_EXTRA is not None:     # mapa sin código postal (entregable suelto): se calcula
        il, io = P["campos"].index("lat"), P["campos"].index("lon")
        z = CPS.asignar([f[il] for f in P["filas"]], [f[io] for f in P["filas"]], DENUE_EXTRA).codigo_postal.tolist()
        P["campos"].append("zip"); [f.append(v) for f, v in zip(P["filas"], z)]
    esc = {c["id"]: c.get("escala") for c in D["capas"]}
    dist_nse = next(c["distribucion"] for c in D["capas"] if c["id"] == "hex_N")
    tablas = tablas_res(D)
    tiene_wh = "wh" in D["hex"]

    FICHA = [["§Código postal (zipcode, DENUE)", ""], ["Código postal que predomina", "zip"], ["Códigos postales (PDV)", "zips"],
             ["§Hogares (INEGI, 2025)", ""], ["Hogares", "hog_h"], ["Nivel NSE de los hogares (1 = D/E … 6 = A/B)", "N_h"],
             ["% bajo (D+, D/E)", "p_bajo"], ["% medio (C, C−)", "p_medio"], ["% alto (C+, A/B)", "p_alto"],
             ["Índice Censo por manzana (0-100)", "censo_h"],
             ["§Clúster NSE (buffer de 300 m)", ""], ["Clústeres NSE con PDV", "n_cl"], ["Nivel NSE INEGI (1-6)", "N_cl"],
             ["Índice AltScore (0-100)", "alt"], ["Nivel NSE final (INEGI + AltScore)", "Nf"],
             ["§Puntos de venta Tradicional", ""], ["PDV", "n_pdv"], ["PDV con letras", "n_letra"]] + \
            [[f"{k} · {ACCION[k]}", f"n_{k}"] for k in LET] + \
            [["Letras que predominan", "letras"], ["Acción Golden Stores", "accion"], ["Prioridad que predomina", "prioridad"]] + \
            [[k, f"n_{k}"] for k in PRI] + \
            ([["§Whisp (hexágono res 6)", ""], ["Share Coca-Cola (%)", "share"], ["Share alto / bajo (vs área)", "share_clase"]] if tiene_wh else [])
    EXTRA = [["Código postal", "zip"], ["Hogares", "hog_h"], ["PDV Tradicional", "n_pdv"], ["Letras que predominan", "letras"], ["Prioridad que predomina", "prioridad"]]
    # paso → (campo, escala, nombre, distribución) en cada resolución
    VISTAS = {
        "1.1": ("N_h", esc["hex_N"], "NSE de los hogares del hexágono (media 1-6)", dist_nse),
        "1.2": ("censo_h", esc["hex_censo"], "Índice del Censo por manzana (0-100)", None),
        "1.3": ("N_cl", esc["cl_N"], "Clúster NSE: nivel INEGI de su buffer (media 1-6)", None),
        "1.4": ("alt", esc["cl_alt"], "Clúster NSE: índice AltScore de su buffer (0-100)", None),
        "1.5": ("Nf", esc["cl_Nf"], "Clúster NSE: nivel final INEGI + AltScore (media 1-6)", None),
        "3.": ("letras", esc["pdv_cl"], "Letras finales que predominan entre sus PDV",
               [[esc["pdv_cl"]["etiquetas"][k], f"p_{k}", esc["pdv_cl"]["colores"][k]] for k in LET]),
        "4.": ("prioridad", esc["pdv_prio"], "Prioridad que predomina entre sus PDV", None),
        "5.": ("zip", ESC_ZIP(tablas), "Código postal que predomina entre sus PDV", None),
    }
    capas_nuevas, por_paso = [], {k: [] for k in VISTAS}
    for r in RES:
        D["hex"][f"r{r}"] = tablas[r]
        for paso, (campo, escala, nombre, dist) in VISTAS.items():
            cid = f"{paso.strip('.').replace('.', '_')}_r{r}"
            c = {"id": cid, "res": r, "nombre": f"Res {r} · {nombre}", "tipo": "hex", "fuente": f"r{r}", "campo": campo, "escala": escala,
                 "borde": True, "opacidad": 0.75, "extra": [e for e in EXTRA if e[1] != campo], "ficha": FICHA,
                 "distribucion_ficha": dist_nse, "titulo_distribucion": "NSE de los hogares del hexágono"}
            if dist:
                c["distribucion"] = dist
            if campo == "zip":
                c.update(buscar_campo="zip", buscar_etiqueta="CP", opacidad=0.6)
            capas_nuevas.append(c); por_paso[paso].append(cid)
    nota = (" <b>Versión KO:</b> sin puntos; todo se ve por hexágono en la resolución que elijas arriba (res 9 = clúster NSE ≈ 0.1 km², "
            "res 8 ≈ 0.7 km², res 7 ≈ 5.2 km²). Clic en un hexágono: todos sus datos; también se busca por su id.")
    pasos = []
    for s in D["pasos"]:
        paso = next((k for k in VISTAS if s["titulo"].startswith(k + " ") or s["titulo"].startswith(k)), None)
        if paso is None:
            continue                                              # pasos que solo tenían puntos (Letra 2, letras base)
        s["capas"] = (["wh_share"] if paso == "4." and tiene_wh else []) + por_paso[paso]
        if paso == "1.1":
            s["titulo"] = "1.1 NSE de los hogares del hexágono"
        s["que"] += nota
        pasos.append(s)
    pasos.append({"titulo": "5. Código postal (zipcode)", "capas": por_paso["5."],
                  "que": "Cada hexágono se pinta con el código postal que más se repite entre sus PDV (color por código; el número está en el "
                         "tooltip y en la ficha, con todos los códigos del hexágono). Escribe un código postal en Buscar para ver todos sus "
                         "hexágonos en la resolución elegida." + nota,
                  "porque": "El CP del cliente no trae código postal: cada PDV toma el del establecimiento DENUE de INEGI más cercano (≤ 500 m).",
                  "formula": "código postal del hexágono = el que más se repite entre sus PDV (empate → el menor)"})
    D["pasos"] = pasos
    D["capas"] = [c for c in D["capas"] if c.get("siempre") or c["id"] == "wh_share"] + capas_nuevas
    # el KO no lleva puntos ni las fuentes finas: solo las tablas por resolución, Whisp y el límite municipal
    D["hex"] = {k: v for k, v in D["hex"].items() if k in ("wh",) or k.startswith("r")}
    D["poligonos"] = {k: v for k, v in D["poligonos"].items() if any(c.get("fuente") == k for c in D["capas"])}
    D["datasets"], D["popup"], D["cluster"], D["filtro"], D["dist_popup"] = {}, {}, None, None, {}
    D["resoluciones"] = {"etiqueta": "Hexágono (resolución H3) · aplica a todos los pasos", "defecto": 9,
                         "valores": [{"valor": r, "etiqueta": f"res {r} · {AREA[r]}", "n": len(tablas[r]["filas"])} for r in RES]}
    D["buscar_placeholder"] = "ID del hexágono o código postal · Enter abre el primero"
    if not D["titulo"].endswith(" · KO"):
        D["titulo"] += " · KO"
    dst.write_text(MP.pagina(D), encoding="utf-8")
    print(f"[{CIUDAD}] mapa KO: hexágonos res 9 {len(tablas[9]['filas']):,} · res 8 {len(tablas[8]['filas']):,} · "
          f"res 7 {len(tablas[7]['filas']):,} · {len(pasos)} pasos, todos por resolución → {dst.name}")


# ---------------------------------------------------------------- Excel reducido
COLORES = {"Punto de venta (CP)": X.NEGRO, "AltScore (buffer del clúster NSE)": "4A3AA7", "Datos gubernamentales (INEGI)": "1F3A5F",
           "Ventas (venta media y potencial del CP · Rappi del buffer · rangos y score)": "1E5B45",
           "Clúster NSE (hexágono H3 + 300 m)": "333333", "Letra 1": X.NEGRO, "Letra 2": X.NEGRO, "Clasificación": "333333"}   # = COLORES del 06
FMT = {"Hogares del clúster (buffer de 300 m)": "#,##0", "Venta media (cajas/mes)": "#,##0.00", "CP · potencial (cajas/mes)": "#,##0.00",
       "Rappi · pedidos/mes totales del buffer": "#,##0.00", "Rappi (pedidos/mes medios de sus tiendas)": "#,##0.00",
       "Rappi · venta del buffer (MXN/mes, referencia)": "#,##0", "Rango venta (0 = mejor)": "0.000", "Rango CP": "0.000", "Rango Rappi": "0.000",
       "Score sin Rappi": "0.000", "Score": "0.000", "Distancia al centro del clúster (m)": "#,##0", "% ABC+ del clúster": "0.0"}  # = FMT del 06
ANCHOS = {"#": 5, "bloque": 34, "columna": 40, "definición": 80, "de dónde sale / fórmula": 60, "unidad o valores": 14}
ZIP = "Código postal (zipcode)"


def excel_ko(src: Path, dst: Path):
    wb0 = openpyxl.load_workbook(src)
    ws = wb0["Letras por PDV"]
    grupos, subs, fuentes_, g = [], [], [], None
    for j in range(1, ws.max_column + 1):
        g = ws.cell(4, j).value or g
        grupos.append(g); subs.append(ws.cell(5, j).value); fuentes_.append(ws.cell(6, j).value)
    cols = list(zip(grupos, subs))
    df = pd.DataFrame(list(ws.iter_rows(min_row=7, values_only=True)), columns=pd.MultiIndex.from_tuples(cols))
    col = lambda sub: next(c for c in cols if c[1] == sub)
    tiene = lambda sub: any(c[1] == sub for c in cols)
    # solo filas con valor (usuario, 2026-10-23): fuera los PDV sin letras y, si el archivo trae Whisp, los sin share Whisp
    con_letras = df[col("Letras (final)")].isin(LET)
    share = pd.to_numeric(df[col("Share KO Whisp (hexágono res 6)")], errors="coerce") if tiene("Share KO Whisp (hexágono res 6)") else None
    con_whisp = share.notna() if share is not None and share.notna().any() else pd.Series(True, index=df.index)
    fuera = (len(df), int((~con_letras).sum()), int((con_letras & ~con_whisp).sum()))
    df = df[con_letras & con_whisp].reset_index(drop=True)
    zip_calc = not tiene(ZIP) and DENUE_EXTRA is not None             # entregable suelto sin código postal: se calcula
    if zip_calc:
        cz = (col("pos_id")[0], ZIP)
        df[cz] = CPS.asignar(df[col("Latitud")], df[col("Longitud")], DENUE_EXTRA).codigo_postal.to_numpy()
        cols.append(cz); fuentes_.append(f"INEGI DENUE · cod_postal del establecimiento más cercano (≤ {CPS.RADIO_MAX_M} m)")

    CLU, POS, SUB = col("Clúster NSE (H3 res 9)"), col("pos_id"), col("Subcanal")
    LAT, LON = col("Latitud"), col("Longitud")
    HEX = {r: (CLU[0], f"Hexágono (H3 res {r})") for r in RES}
    H = "Clasificación"
    QUITAR = {"Nombre", "Latitud", "Longitud"}
    PRIMERO = ([col(ZIP)] if tiene(ZIP) else []) + [POS, SUB]          # zipcode junto a los hexágonos (usuario, 2026-10-22)
    RED = pd.DataFrame({HEX[r]: [hex_res(c9, la, lo, r) for c9, la, lo in zip(df[CLU], df[LAT], df[LON])] for r in RES})
    for c in PRIMERO + [c for c in cols if c[1] not in QUITAR and c != CLU and c not in PRIMERO]:
        RED[c] = df[c].to_numpy()
    for r in RES:
        let_h = df[col("Letras (final)")].groupby(RED[HEX[r]].to_numpy()).agg(lambda s: moda(list(s), LET))
        pri_h = df[col("Prioridad")].groupby(RED[HEX[r]].to_numpy()).agg(lambda s: moda(list(s), PRI))
        RED[(H, f"Letras del hexágono res {r}")] = RED[HEX[r]].map(let_h).fillna(SIN).to_numpy()
        RED[(H, f"Acción del hexágono res {r}")] = RED[(H, f"Letras del hexágono res {r}")].map(ACCION).fillna(SIN).to_numpy()
        RED[(H, f"Prioridad del hexágono res {r}")] = RED[HEX[r]].map(pri_h).fillna(SIN).to_numpy()
        RED[HEX[r]] = RED[HEX[r]].fillna("sin hexágono (sin coordenadas)")
    RED.columns = pd.MultiIndex.from_tuples(list(RED.columns))

    fuentes = dict(zip(cols, fuentes_))
    fuentes[HEX[9]] = fuentes.pop(CLU)
    for r in (8, 7):
        fuentes[HEX[r]] = f"hexágono H3 res {r} ({AREA[r]}) que contiene al PDV (latitud y longitud del CP)"
    for r in RES:
        fuentes[(H, f"Letras del hexágono res {r}")] = f"la letra final que más se repite entre los PDV del hexágono res {r} (empate → HH, HL, LH, LL)"
        fuentes[(H, f"Acción del hexágono res {r}")] = "HH Atacar · HL Bloquear · LH Fortalecer · LL Mantener"
        fuentes[(H, f"Prioridad del hexágono res {r}")] = f"la prioridad que más se repite entre los PDV del hexágono res {r} (empate → P1 … P4)"
    wb = X.libro()
    X.hoja_tabla_2niveles(wb, "Letras por PDV", RED, titulo="Letras por punto de venta · canal Tradicional (reducido · KO)",
                          nota="Solo PDV con letras (y con share Whisp si el área tiene Whisp). Sin coordenadas: cada PDV lleva su hexágono H3 en res 9 (clúster NSE), res 8 y res 7"
                               + (" y su código postal" if tiene(ZIP) else "") + ". Al final, la letra, la acción y la prioridad que "
                               "predominan en cada hexágono (las del mapa KO). La 3.ª fila dice de dónde sale cada columna; la hoja Diccionario define cada una.",
                          formatos=FMT, anchos={ZIP: 14}, colores_grupo=COLORES, fijar_cols=3 + tiene(ZIP), fuentes=fuentes)

    d0 = wb0["Diccionario"]
    dic = pd.DataFrame(list(d0.iter_rows(min_row=5, values_only=True)), columns=[c.value for c in d0[4]])
    if zip_calc:
        dic = pd.concat([dic, pd.DataFrame([{"bloque": col("pos_id")[0], "columna": ZIP, "definición": "Código postal (zipcode) del PDV: el del "
                         f"establecimiento DENUE de INEGI más cercano (≤ {CPS.RADIO_MAX_M} m).", "de dónde sale / fórmula": fuentes_[-1],
                         "unidad o valores": "5 dígitos"}])], ignore_index=True)
    fila = lambda n: dic[dic["columna"] == n]
    filas_hex = []
    for r in RES:
        h = fila(CLU[1]).copy()
        h["bloque"], h["columna"] = HEX[r]
        h["definición"] = ("Hexágono H3 res 9 (≈ 0.1 km²) del clúster NSE donde está el PDV; reemplaza a la latitud y la longitud. "
                           "Todos sus PDV comparten la Letra 1." if r == 9 else
                           f"Hexágono H3 res {r} ({AREA[r]}) que contiene al PDV: agrupa más PDV que el clúster NSE.")
        h["de dónde sale / fórmula"], h["unidad o valores"] = fuentes[HEX[r]], "id H3"
        filas_hex.append(h)
    resto = dic[~dic["columna"].isin(QUITAR | {CLU[1]} | {c[1] for c in PRIMERO})]
    extra = []
    for r in RES:
        extra += [
            {"bloque": H, "columna": f"Letras del hexágono res {r}", "definición": f"Letras finales que más se repiten entre los PDV del hexágono res {r} "
             "(las que pinta el mapa KO en esa resolución); empate → la mejor.", "de dónde sale / fórmula": fuentes[(H, f"Letras del hexágono res {r}")],
             "unidad o valores": "HH · HL · LH · LL"},
            {"bloque": H, "columna": f"Acción del hexágono res {r}", "definición": f"Acción Golden Stores de las letras del hexágono res {r}.",
             "de dónde sale / fórmula": fuentes[(H, f"Acción del hexágono res {r}")], "unidad o valores": "—"},
            {"bloque": H, "columna": f"Prioridad del hexágono res {r}", "definición": f"Prioridad que más se repite entre los PDV del hexágono res {r} "
             "(la que pinta el mapa KO en esa resolución); empate → la mejor.", "de dónde sale / fórmula": fuentes[(H, f"Prioridad del hexágono res {r}")],
             "unidad o valores": "P1 · P2 · P3 · P4"}]
    dic = pd.concat(filas_hex + [fila(c[1]) for c in PRIMERO] + [resto, pd.DataFrame(extra)], ignore_index=True)
    dic["#"] = range(1, len(dic) + 1)
    assert list(dic["columna"]) == [c[1] for c in RED.columns], "el Diccionario no sigue el orden de las columnas"
    X.hoja_tabla(wb, "Diccionario", dic, titulo="Diccionario de datos · Letras por PDV (reducido · KO)",
                 nota="Una fila por columna de la hoja Letras por PDV, en el mismo orden.", anchos=ANCHOS)
    wb.save(dst)
    igual = {r: (RED[(H, f"Letras del hexágono res {r}")] == RED[col("Letras (final)")]).mean() for r in RES}
    print(f"[{CIUDAD}] Excel KO: de {fuera[0]:,} PDV fuera {fuera[1]:,} sin letras y {fuera[2]:,} sin share Whisp · "
          f"{len(RED):,} PDV · {RED.shape[1]} columnas · hexágonos " +
          " · ".join(f"res {r} {RED[HEX[r]].nunique():,} (letra = la del PDV {igual[r]:.0%})" for r in RES) +
          f" · código postal {'sí' if tiene(ZIP) else 'no (el reducido no lo trae)'} → {dst.name}")


if __name__ == "__main__":
    if sys.argv[1:2] == ["--archivos"]:
        m, r, m_ko, r_ko = map(Path, sys.argv[2:6])
        if "--estados" in sys.argv:                                  # p. ej. --estados 09,15 (CDMX y Edomex)
            DENUE_EXTRA = CPS.denue_estados(sys.argv[sys.argv.index("--estados") + 1].split(","))
        CIUDAD = m.stem
        mapa_ko(m, m_ko)
        excel_ko(r, r_ko)
        sys.exit()
    if CARPETA:                      # una carpeta cliente: el mapa y el reducido NORMALES que haya ahí → sus _KO al lado
        normal = lambda ps: sorted(p for p in ps if not p.stem.upper().endswith("_KO") and not p.name.startswith("~$"))
        mapas_ = normal(CARPETA.glob("06_mapa_qa_letras_*.html")) or normal(CARPETA.glob("*mapa*.html"))
        reds = normal(CARPETA.glob("*reducido*.xlsx"))
        if len(mapas_) != 1 or len(reds) != 1:
            sys.exit(f"[{CIUDAD}] en {CARPETA} debe haber 1 mapa y 1 Excel reducido normales (no _KO); hay mapas {[p.name for p in mapas_]} "
                     f"y reducidos {[p.name for p in reds]}")
        if "--estados" in sys.argv:
            DENUE_EXTRA = CPS.denue_estados(sys.argv[sys.argv.index("--estados") + 1].split(","))
        elif CONFIGURADA:            # ciudad configurada: su DENUE (solo se usa si el mapa o el Excel no traen código postal)
            DENUE_EXTRA = CPS.denue_zm()
        else:
            print(f"[{CIUDAD}] ciudad sin configurar y sin --estados: si el Excel no trae código postal, el KO sale sin él")
        mapa_ko(mapas_[0], mapas_[0].with_name(mapas_[0].stem + "_KO.html"))
        excel_ko(reds[0], reds[0].with_name(reds[0].stem + "_KO.xlsx"))
        sys.exit()
    base = f"{C.CLIENTE}_{C.SLUG}"
    mapa = C.OUT_CLIENTE / f"06_mapa_qa_letras_{base}.html"          # los mapas van en cliente/ (usuario, 2026-10-23)
    red = C.OUT_CLIENTE / f"06_letras_nse_ventas_tradicional_reducido_{base}.xlsx"
    for f in (mapa, red):
        if not f.exists():
            sys.exit(f"[{CIUDAD}] falta {f}: correr antes el 06 (uv run python src/correr.py --pasos 06 {CIUDAD})")
    mapa_ko(mapa, mapa.with_name(mapa.stem + "_KO.html"))
    excel_ko(red, red.with_name(red.stem + "_KO.xlsx"))
