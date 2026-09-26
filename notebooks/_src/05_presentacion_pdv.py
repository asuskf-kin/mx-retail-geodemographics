# %% [markdown]
# # 05 · Golden Stores — canal Tradicional (Bepensa · Sueros · ZM Mérida) — último paso
#
# Replica la **metodología Golden Stores** (NielsenIQ | Spectra, "Golden Stores Sueros Nacional Pharma FY'23") sobre
# los puntos de venta (PDV) del canal **Tradicional** de Bepensa, con el estilo visual de Kin (`src/deck_kin.py`).
# **La metodología no se modifica**; solo se sustituyen las fuentes por las disponibles:
#
# | Golden Stores (NIQ) | Aquí |
# |---|---|
# | Ventas de la tienda (valor / volumen del periodo, Scantrack) | Cajas de sueros por mes de vida del PDV (ventas Bepensa, validadas en el 00) |
# | Demanda potencial con base en su área transaccional (Spectra) | Hogares 2025 a ≤ 300 m del PDV: manzanas urbanas (Censo 2020 + microsimulación) y localidades rurales (ITER 2020) repartidas por área, actualizadas a 2025 con la Encuesta Intercensal 2025 (notebooks 01 y 04) |
# | Mercado / target | Canal Tradicional, ZM Mérida, sueros |
#
# **1. Cuatro clústeres** (índice 100 = promedio del mercado; se calcula dentro de cada subcanal porque el 00 mostró que el
# subcanal cambia la escala de venta):
#
# | Clúster | Demanda | Venta | Acción | Accionables |
# |---|---|---|---|---|
# | **HH** (Golden Stores) | alta | alta | **Atacar** | más SOS, exhibiciones, special packs, ítems exclusivos, reducciones de precio, comunicación, promoción |
# | **HL** | alta | baja | **Bloquear** | evitar desarrollo de competencia, estrategia reactiva, seguir estrategias de precio, cuidar anaquel (SOS) |
# | **LH** | baja | alta | **Fortalecer** | incrementar exhibiciones, cuidar anaquel (SOS), considerar special packs |
# | **LL** | baja | baja | **Mantener** | mantener distribución, mantener anaquel |
#
# **2. Panorama por clúster:** # tiendas (%), % mix de ventas, índice de ventas (promedio del clúster / promedio total × 100),
# perfil de hogares (tamaño, niños, edad del jefe, NSE: % e índice vs ZM), mix por municipio y subcanal, clase CP.
#
# **3. Semáforos por clúster** (individualizados, con desviación estándar dentro de cada clúster; venta y demanda en log
# por su cola pesada, así los dos ejes quedan en la misma escala): **verde** = venta y demanda sobre el promedio del clúster; **rojo** = ambas bajo el promedio;
# **amarillo** = alrededor de los promedios (|z| ≤ 0.5 en ambas, #2) o subdesarrollo en una (#1 alta venta / demanda baja,
# #3 demanda alta / venta menor).
#
# **4. Priorización P1:** capitalizar los clústeres de alta venta **HH y LH** con semáforo **verde y amarillo**.
#
# **Customer Potential (CP):** potencial *futuro* de cada cliente frente a su comparable; se usa como variable de
# **clasificación** (clase Very High → Low), no como criterio decisivo.
#
# **Controles de calidad (QA metodológico, se recalculan en cada corrida):**
# * **Frontera:** tiendas a ±10 puntos del índice 100 en alguna variable (su clúster cambia con poco ruido): marcadas en el Excel.
# * **Actividad:** % de tiendas activas / en riesgo / inactivas por clúster. Una HL que dejó de comprar no se "bloquea": primero se reactiva.
# * **Brecha HL vs HH:** sale de cómo se define el clúster (se comprueba con venta al azar), así que no se presenta como oportunidad.
# * **Demanda vs venta:** correlación de Spearman de hogares con la venta dentro del canal (la demanda clasifica, no pronostica).
# * **Excluidas:** tiendas con venta y sin hogares a ≤ 300 m (zonas comerciales): van en hoja aparte, no desaparecen.
# * **Piloto:** sorteo por zonas H3 res 7 (≈ 5 km²) para que las áreas de 300 m del grupo control no se mezclen con las
#   tratadas, y efecto mínimo detectable (MDE) con el efecto de diseño por zona. Se compara piloto vs control, nunca
#   antes vs después (las tiendas elegidas por venta alta bajan solas: regresión a la media).
#
# **Salidas:** `outputs/<ciudad>/05_golden_stores_<canal>_bepensa_zm_<ciudad>.pptx` y `.xlsx` (entregable por tienda).
# **Reproducir:** orden 00 → 01 → 02 → 03 → 04 → 05 (`python src/correr.py --pasos 00,01,02,03,04,05 merida`).

# %%
import os
import sys
from pathlib import Path

os.environ.setdefault("CIUDAD", "merida")
BASE = next(p for p in [Path.cwd(), *Path.cwd().parents] if (p / "src" / "config.py").exists())
sys.path.insert(0, str(BASE / "src"))
import numpy as np
import pandas as pd
from IPython.display import display

import config as C
import eda
from descargas import requisitos

try:
    import deck_kin as K
except ImportError as e:
    raise ImportError("Falta python-pptx: python -m pip install python-pptx==1.0.2") from e

CANALES_DECK = ["Tradicional"]          # canal(es) sobre los que se replica Golden Stores (agregar "Moderno" si se requiere)
assert C.CIUDAD == C.CLIENTE_CIUDAD, f"Los datos de {C.CLIENTE} son de {C.CLIENTE_CIUDAD}; CIUDAD={C.CIUDAD}"
F00 = C.OUT / f"00_validacion_ventas_cp_{C.CLIENTE}_{C.SLUG}.xlsx"
F04 = C.PROC / f"pdv_{C.CLIENTE}_hexagonos_{C.SLUG}.pkl"
requisitos({F00: "00_ventas_cp_validacion", F04: "04_hexagonos_pdv", C.PROC / f"pdv_{C.CLIENTE}_{C.SLUG}.parquet": "00_ventas_cp_validacion",
            C.PROC / f"nse_ageb_{C.SLUG}.parquet": "01_socioeconomico_merida",
            C.PROC / f"robustez_{C.CLIENTE}_{C.SLUG}.parquet": "04_hexagonos_pdv"})
pd.set_option("display.width", 220, "display.max_columns", 30, "display.float_format", "{:,.2f}".format)

CLUSTERS = ["HH", "HL", "LH", "LL"]
ACCION = {"HH": "Atacar", "HL": "Bloquear", "LH": "Fortalecer", "LL": "Mantener"}
NOMBRE = {"HH": "Alta demanda / Alta venta", "HL": "Alta demanda / Baja venta", "LH": "Baja demanda / Alta venta", "LL": "Baja demanda / Baja venta"}
ACCIONABLES = {"HH": ["Más SOS", "Exhibiciones", "Special packs", "Ítems exclusivos", "Reducciones de precio", "Comunicación", "Promoción"],
               "HL": ["Evitar desarrollo de competencia", "Estrategia reactiva", "Seguir estrategias de precio", "Cuidar anaquel (SOS)"],
               "LH": ["Incrementar exhibiciones", "Cuidar anaquel (SOS)", "Considerar special packs"],
               "LL": ["Mantener distribución", "Mantener anaquel"]}
COLOR = {"HH": K.NEGRO, "HL": K.GRIS, "LH": K.VERDE, "LL": K.GRIS_MEDIO}
SEMAFORO = {"verde": "22C55E", "amarillo": "EDA100", "rojo": "E34948"}
ORDEN_CP = ["Very High", "High", "Moderate", "Low"]

# %% [markdown]
# ## 1. Datos

# %%
tabla = pd.read_pickle(F04)
ids = tabla[""].copy()
ids["Subcanal"] = ids.Subcanal.str.replace("ABARRROTES", "ABARROTES")      # error de captura del cliente (solo para mostrar)
hh = pd.DataFrame({c: tabla[(c, "HHs")] for c in C.CATEGORIAS})
ageb = pd.read_parquet(C.PROC / f"nse_ageb_{C.SLUG}.parquet")
REF = pd.Series(ageb[[f"HHs_{c}" for c in C.CATEGORIAS]].sum().values / ageb["Total HHs"].sum() * 100, index=C.CATEGORIAS)
h00 = pd.read_excel(F00, sheet_name=None)
rob = pd.read_parquet(C.PROC / f"robustez_{C.CLIENTE}_{C.SLUG}.parquet")
rob["subcanal"] = rob.subcanal.str.replace("ABARRROTES", "ABARROTES")
pdv0 = pd.read_parquet(C.PROC / f"pdv_{C.CLIENTE}_{C.SLUG}.parquet").set_index("pos_id_cp")
d = ids.join(hh)
d["clase_cp"] = d.pos_id_cp.map(pdv0.PotentialQualitative_TotalPortafolio)
print(f"PDV: {len(d):,} | con venta: {int(d['Con venta'].sum()):,} | canales: {d.Canal.value_counts().to_dict()}")

# %% [markdown]
# ## 2. Clústeres, panorama, perfil y semáforos (por canal)

# %%
def golden_stores(canal: str):
    """Clústeres HH/HL/LH/LL (índice 100 = promedio del subcanal), panorama, perfil de hogares y semáforos."""
    b = d[(d.Canal == canal) & d["Con venta"].astype(bool) & (d["HHs en el área"] > 0)].copy()
    g = b.groupby("Subcanal")
    b["indice_demanda"] = b["HHs en el área"] / g["HHs en el área"].transform("mean") * 100
    b["indice_venta"] = b["Cajas/mes (vida)"] / g["Cajas/mes (vida)"].transform("mean") * 100
    b["cluster"] = np.where(b.indice_demanda >= 100, "H", "L") + np.where(b.indice_venta >= 100, "H", "L")
    b["accion"] = b.cluster.map(ACCION)
    # semáforos individualizados por clúster (desviación estándar dentro del clúster)
    gc = b.groupby("cluster")
    lv, ld = np.log1p(b["Cajas/mes (vida)"]), np.log1p(b["HHs en el área"])          # misma escala (log) en los dos ejes
    zv = (lv - lv.groupby(b.cluster).transform("mean")) / lv.groupby(b.cluster).transform("std")
    zd = (ld - ld.groupby(b.cluster).transform("mean")) / ld.groupby(b.cluster).transform("std")
    b["z_venta"], b["z_demanda"] = zv, zd
    alrededor = (zv.abs() <= 0.5) & (zd.abs() <= 0.5)
    b["semaforo"] = np.select([alrededor, (zv > 0) & (zd > 0), (zv < 0) & (zd < 0)], ["amarillo", "verde", "rojo"], "amarillo")
    b["semaforo_detalle"] = np.select([alrededor, (zv > 0) & (zd > 0), (zv < 0) & (zd < 0), zv > 0],
                                      ["#2 alrededor del promedio", "alta venta y alta demanda", "baja venta y baja demanda",
                                       "#1 alta venta / demanda baja"], "#3 demanda alta / venta menor")
    b["frontera"] = ((b.indice_demanda - 100).abs() <= 10) | ((b.indice_venta - 100).abs() <= 10)
    b["alta_reciente"] = b.pos_id_cp.map(pdv0.meses_vida) <= 3
    b["brecha_relativa"] = b.pos_id_cp.map(pdv0.PotentialEstimatedToCover_TotalPortafolio)
    media = b["Cajas/mes (vida)"].mean()
    pan = (b.groupby("cluster").agg(tiendas=("pos_id_cp", "size"), cajas=("Cajas/mes (vida)", "sum"), media=("Cajas/mes (vida)", "mean"),
                                     hogares=("HHs en el área", "median"), activos=("Estado actividad", lambda x: (x == "activo").mean() * 100))
           .reindex(CLUSTERS))
    pan["% tiendas"] = pan.tiendas / pan.tiendas.sum() * 100
    pan["% mix ventas"] = pan.cajas / pan.cajas.sum() * 100
    pan["Index ventas"] = pan.media / media * 100
    perfil = b.groupby("cluster")[C.CATEGORIAS].sum().reindex(CLUSTERS)
    pct = perfil.div(perfil[C.GRUPOS["Tamaño hogar"]].sum(axis=1), axis=0) * 100
    for grupo, cols in C.GRUPOS.items():               # cada grupo suma 100
        pct[cols] = pct[cols].div(pct[cols].sum(axis=1), axis=0) * 100
    indice = pct.div(REF, axis=1) * 100
    sem = pd.crosstab(b.cluster, b.semaforo).reindex(index=CLUSTERS, columns=["verde", "amarillo", "rojo"]).fillna(0).astype(int)
    return b, pan, pct, indice, sem


res = {c: golden_stores(c) for c in CANALES_DECK}
for c, (b, pan, pct, indice, sem) in res.items():
    print(f"== {c}: {len(b):,} tiendas (target)")
    display(pan[["tiendas", "% tiendas", "% mix ventas", "Index ventas", "hogares", "activos"]].round(1))
    display(indice.round(0))
    display(sem)

# %% [markdown]
# ### 2b. Controles de calidad (QA metodológico)

# %%
from scipy import stats
EXCL, QA = {}, {}
rng_qa = np.random.default_rng(eda.SEMILLA)
for c, (b, pan, pct, indice, sem) in res.items():
    V, D = "Cajas/mes (vida)", "HHs en el área"
    EXCL[c] = d[(d.Canal == c) & d["Con venta"].astype(bool) & ~(d[D] > 0)].copy()
    act = (pd.crosstab(b.cluster, b["Estado actividad"], normalize="index") * 100).reindex(CLUSTERS)
    # brecha HL vs HH con la venta permutada dentro de subcanal (sin relación con la demanda): ¿es por construcción?
    gaps = []
    for _ in range(200):
        v = b.groupby("Subcanal")[V].transform(lambda x: pd.Series(rng_qa.permutation(x.to_numpy()), index=x.index))
        iv = v / v.groupby(b.Subcanal).transform("mean") * 100
        cl = np.where(b.indice_demanda >= 100, "H", "L") + np.where(iv >= 100, "H", "L")
        m = v.groupby(cl).mean()
        gaps.append(1 - m["HL"] / m["HH"])
    rho = stats.spearmanr(b[D], b[V])
    p1q = b[b.cluster.isin(["HH", "LH"]) & b.semaforo.isin(["verde", "amarillo"])]
    QA[c] = dict(actividad=act, brecha_real=1 - pan.loc["HL", "media"] / pan.loc["HH", "media"], brecha_azar=float(np.mean(gaps)),
                 rho_demanda=rho[0], p_demanda=rho[1], frontera=b.frontera.mean(), p1_de_hhlh=len(p1q) / b.cluster.isin(["HH", "LH"]).sum(),
                 excl_n=len(EXCL[c]), excl_mix=EXCL[c][V].sum() / (EXCL[c][V].sum() + b[V].sum()),
                 hl_no_activas=int((b[b.cluster == "HL"]["Estado actividad"] != "activo").sum()),
                 med_act={q: b[(b.cluster == q) & (b["Estado actividad"] == "activo")][V].median() for q in CLUSTERS})
    q = QA[c]
    print(f"== {c} · QA")
    print(f"Brecha de venta HL vs HH: {q['brecha_real']:.0%} real vs {q['brecha_azar']:.0%} con venta al azar → sale de la definición del clúster; no es oportunidad.")
    print(f"Hogares a 300 m vs venta: Spearman {q['rho_demanda']:+.3f} (p = {q['p_demanda']:.1g}) → la demanda clasifica, no pronostica.")
    print(f"Frontera (±10 puntos del índice 100): {q['frontera']:.1%} de las tiendas | P1 conserva {q['p1_de_hhlh']:.0%} de HH + LH.")
    print(f"HL no activas (en riesgo o inactivas): {q['hl_no_activas']:,} | mediana cajas/mes de activas: "
          + " · ".join(f"{k} {v:.1f}" for k, v in q["med_act"].items()))
    print(f"Excluidas por no tener hogares a ≤ {C.RADIO_PDV_M} m: {q['excl_n']:,} tiendas ({q['excl_mix']:.1%} de la venta del canal) → hoja aparte en el Excel.")
    display(act.round(1))

# %% [markdown]
# ### 2c. Diseño del piloto: sorteo por zonas y efecto mínimo detectable
#
# Las áreas de 300 m se traslapan, así que tiendas tratadas y de control vecinas comparten clientes (contaminación). Se sortean
# **zonas H3 res 7** completas (≈ 5 km²), emparejadas por número de tiendas, y el MDE incluye el efecto de diseño
# (1 + (m − 1)·ICC, con el ICC del log de la venta entre zonas). Sin venta mensual no hay línea base: el MDE "sin línea base"
# es el vigente; el "con línea base" supone correlación 0.8 entre periodos (supuesto: se confirma con la venta mensual).

# %%
import h3
Z = 1.959964 + 0.841621                                   # α = 5% bilateral, potencia 80%


def mde(y, zona, n_t, n_c, rho_base=0.0):
    ly = pd.Series(np.log1p(y.to_numpy(float)))
    g = ly.groupby(zona.to_numpy())
    k, n = g.ngroups, len(ly)
    if k < 2:
        return np.nan, np.nan, np.nan
    m0 = (n - (g.size() ** 2).sum() / n) / (k - 1)
    msb = (g.size() * (g.mean() - ly.mean()) ** 2).sum() / (k - 1)
    msw = ((ly - g.transform("mean")) ** 2).sum() / max(n - k, 1)
    icc = max(0.0, (msb - msw) / (msb + (m0 - 1) * msw))
    deff = 1 + (n / k - 1) * icc
    return float(np.expm1(Z * ly.std() * np.sqrt(deff * (1 / n_t + 1 / n_c) * (1 - rho_base ** 2)))), icc, deff


def sorteo(u):
    u = u.assign(zona=[h3.latlng_to_cell(a, o, 7) for a, o in zip(u.Latitud, u.Longitud)])
    tam = u.zona.value_counts()
    rg = np.random.default_rng(eda.SEMILLA)
    grupo = {}
    zonas = list(tam.index)                               # emparejar zonas de tamaño parecido y sortear cuál recibe Atacar
    for i in range(0, len(zonas), 2):
        par = zonas[i:i + 2]
        orden = rg.permutation(len(par))
        for j, z in zip(orden, par):
            grupo[z] = "Piloto" if j == 0 else "Control"
    u["grupo"] = u.zona.map(grupo)
    t, cc = u[u.grupo == "Piloto"], u[u.grupo == "Control"]
    la1, lo1, la2, lo2 = (np.radians(x) for x in (cc.Latitud.to_numpy(), cc.Longitud.to_numpy(), t.Latitud.to_numpy(), t.Longitud.to_numpy()))
    dist = 2 * 6_371_000 * np.arcsin(np.sqrt(np.sin((la1[:, None] - la2) / 2) ** 2
                                             + np.cos(la1[:, None]) * np.cos(la2) * np.sin((lo1[:, None] - lo2) / 2) ** 2))
    contam = float((dist.min(axis=1) < 2 * C.RADIO_PDV_M).mean()) if len(t) and len(cc) else np.nan
    return u, contam


PILOTO = {}
for c, (b, pan, pct, indice, sem) in res.items():
    opciones = {"Golden Stores en verde (HH verde)": b[(b.cluster == "HH") & (b.semaforo == "verde")],
                "Golden Stores (HH verde y amarillo)": b[(b.cluster == "HH") & b.semaforo.isin(["verde", "amarillo"])],
                "P1 (HH y LH en verde y amarillo)": b[b.cluster.isin(["HH", "LH"]) & b.semaforo.isin(["verde", "amarillo"])]}
    filas, asignaciones = [], {}
    for nombre, u in opciones.items():
        u, contam = sorteo(u)
        nt, nc = int((u.grupo == "Piloto").sum()), int((u.grupo == "Control").sum())
        m0_, icc, deff = mde(u["Cajas/mes (vida)"], u.zona, nt, nc)
        m8, _, _ = mde(u["Cajas/mes (vida)"], u.zona, nt, nc, rho_base=0.8)
        filas.append({"universo": nombre, "tiendas": len(u), "zonas H3 r7": u.zona.nunique(), "piloto": nt, "control": nc,
                      "ICC zona": icc, "efecto de diseño": deff, "MDE sin línea base": m0_, "MDE con línea base (ρ=0.8)": m8,
                      "control a < 600 m de una tratada": contam})
        asignaciones[nombre] = u
    tab = pd.DataFrame(filas)
    rec = tab.loc[tab["MDE sin línea base"].idxmin(), "universo"]      # el universo que detecta el efecto más chico
    PILOTO[c] = dict(tabla=tab, recomendado=rec, asignacion=asignaciones[rec], fila=tab.set_index("universo").loc[rec])
    print(f"== {c} · piloto recomendado: {rec}")
    display(tab.round(3))

# %% [markdown]
# ## 3. Presentación Golden Stores (estilo Kin)

# %%
def h(tema):
    r = h00["Hallazgos"]
    r = r[r.tema == tema]
    return r.hallazgo.iloc[0] if len(r) else ""


def perfil_cluster(s, x, y, w, q, b, pct, indice, cats, mostrar_etiquetas):
    """Columna de perfil tipo Golden Stores: barra de % por categoría + índice vs ZM (resaltado si ≥ 105)."""
    n = (b.cluster == q).sum()
    chip = K.caja(s, x + (0.95 if mostrar_etiquetas else 0), y, 0.55, 0.26, relleno=COLOR[q], redondeo=0.2)
    K.texto(s, x + (0.95 if mostrar_etiquetas else 0), y + 0.05, 0.55, 0.18, q, tam=10, negrita=True,
            color=K.LIMA if q == "HH" else K.BLANCO, alinear=K.PP_ALIGN.CENTER)
    K.texto(s, x + (1.55 if mostrar_etiquetas else 0.6), y + 0.06, 1.1, 0.18, f"{n:,} ({n / len(b):.0%})", tam=7.5, color=K.GRIS)
    yy = y + 0.36
    x_bar = x + (0.95 if mostrar_etiquetas else 0)
    ancho_max = w - (0.95 if mostrar_etiquetas else 0) - 0.42
    for grupo, cols in cats:
        if mostrar_etiquetas:
            K.texto(s, x, yy, 0.95, 0.14, grupo, tam=6.5, negrita=True, color=K.NEGRO)
        yy += 0.13
        top = pct.loc[q, cols].idxmax()
        for ccat in cols:
            v, iv = pct.loc[q, ccat], indice.loc[q, ccat]
            if mostrar_etiquetas:
                K.texto(s, x + 0.05, yy - 0.01, 0.9, 0.13, ccat, tam=5.5, color=K.GRIS)
            K.caja(s, x_bar, yy + 0.01, max(ancho_max * v / 60, 0.02), 0.09, relleno=COLOR[q] if ccat == top else K.GRIS_CLARO,
                   forma=K.MSO_SHAPE.RECTANGLE)
            K.texto(s, x_bar + max(ancho_max * v / 60, 0.02) + 0.03, yy - 0.01, 0.35, 0.12, f"{v:.0f}%", tam=5.5, color=K.GRIS)
            fondo = iv >= 105
            if fondo:
                K.caja(s, x + w - 0.36, yy - 0.01, 0.32, 0.13, relleno=K.LIMA, redondeo=0.2)
            K.texto(s, x + w - 0.36, yy - 0.005, 0.32, 0.12, f"{iv:.0f}", tam=6, negrita=fondo, alinear=K.PP_ALIGN.CENTER,
                    color=K.NEGRO if fondo or iv > 95 else K.GRIS_MEDIO)
            yy += 0.125
        yy += 0.04


def ejemplos_cluster(b):
    """Una tienda real por clúster para ilustrar (verde con más venta en HH/LH, más demanda en HL, mediana en LL)."""
    ej = {}
    for q in CLUSTERS:
        g = b[b.cluster == q]
        gv = g[g.semaforo == "verde"] if (g.semaforo == "verde").any() else g
        if q in ("HH", "LH"):
            r = gv.sort_values("Cajas/mes (vida)", ascending=False).iloc[0]
        elif q == "HL":
            r = gv.sort_values("HHs en el área", ascending=False).iloc[0]
        else:
            r = g.iloc[(g["Cajas/mes (vida)"] - g["Cajas/mes (vida)"].median()).abs().argsort().iloc[0]]
        ej[q] = r
    return ej


def deck(canal, b, pan, pct, indice, sem):
    n = len(b)
    hh_, hl_, lh_ = pan.loc["HH"], pan.loc["HL"], pan.loc["LH"]
    sesgo = h00["Sesgo ventas sin ubicar"]
    sin_ubicar = sesgo["% de las cajas"].iloc[0]
    p1 = b[b.cluster.isin(["HH", "LH"]) & b.semaforo.isin(["verde", "amarillo"])]
    p1_n, p1_mix = len(p1), p1["Cajas/mes (vida)"].sum() / b["Cajas/mes (vida)"].sum() * 100
    p1_cajas, total_cajas = p1["Cajas/mes (vida)"].sum(), b["Cajas/mes (vida)"].sum()
    hh_verde = int(sem.loc["HH", "verde"])
    caida_hl = 1 - hl_.media / hh_.media
    comp = b.assign(comp=b["DENUE Moderno en el área"] + b["DENUE Tradicional en el área"]).groupby("cluster").comp.median()
    ej = ejemplos_cluster(b)
    e0 = ej["HH"]
    ab = indice[["A/B", "C+"]].mean(axis=1)
    mm = (pd.crosstab(b.Municipio, b.cluster, normalize="columns") * 100).reindex(columns=CLUSTERS).fillna(0)
    mm = mm.loc[mm.sum(axis=1).sort_values(ascending=False).index]
    mun_top = mm["HH"].idxmax()
    cpx = (pd.crosstab(b.cluster, b.clase_cp, normalize="index") * 100).reindex(index=CLUSTERS, columns=ORDEN_CP).fillna(0)
    sens = {p: p1_cajas * p / 100 for p in (5, 10, 15)}
    fmt_c = lambda v: f"{v:,.0f}"
    qa, pil, excl = QA[canal], PILOTO[canal], EXCL[canal]
    act = qa["actividad"]
    hl_na, hl_n = qa["hl_no_activas"], int(hl_.tiendas)
    rec, fr = pil["recomendado"], pil["fila"]
    n_pil, mde0, mde8 = int(fr.tiendas), fr["MDE sin línea base"], fr["MDE con línea base (ρ=0.8)"]
    rec_corto = rec.split(" (")[0]
    p1_front = int(b[b.pos_id_cp.isin(p1.pos_id_cp)].frontera.sum())

    prs = K.nueva()
    s = K.portada(prs, ["Golden Stores"], f"Sueros · Canal {canal} · {C.ZM_NOMBRE}",
                  f"{p1_n:,} tiendas concentran el {p1_mix:.0f}% de la venta: dónde actuar primero y cómo medirlo.",
                  "Bepensa · Septiembre 2026")
    K.notas(s, "Presentación para decidir en qué tiendas del canal Tradicional actuar primero. La recomendación va en la siguiente lámina.")

    # 1 · BLUF: resumen ejecutivo y la decisión
    s, y = K.lamina(prs, f"Activar {p1_n:,} tiendas P1 asegura el {p1_mix:.0f}% de la venta ubicada de sueros en {canal}.", "Resumen ejecutivo",
                    notas=(f"GUION 60 s: De {n:,} tiendas tradicionales con venta en la {C.ZM_NOMBRE}, {p1_n:,} (HH y LH en semáforo verde y amarillo) "
                           f"hacen el {p1_mix:.0f}% de la venta de sueros. Las {int(hh_.tiendas):,} Golden Stores venden {hh_['Index ventas'] - 100:+.0f}% vs el "
                           f"promedio. Pedimos un piloto de 90 días en {rec_corto} ({n_pil:,} tiendas) sorteado por zonas, que detecta efectos desde "
                           f"{mde0:.0%}; con {hh_verde} tiendas el efecto mínimo detectable sería mayor que el +10% que buscamos. Cada +10% en P1 son "
                           f"{fmt_c(sens[10])} cajas al mes (sensibilidad, no pronóstico). "
                           "Objeción probable: '¿por qué no usar solo el CP?' → El CP mide potencial futuro y lo usamos para clasificar; la venta real y la "
                           "demanda alrededor deciden el clúster."))
    for i, (v, e, dd) in enumerate([(f"{p1_n:,}", "tiendas P1", "HH y LH en verde y amarillo"),
                                    (f"{p1_mix:.0f}%", "de la venta del canal", "en esas tiendas"),
                                    (f"{hh_['Index ventas']:.0f}", "índice de venta HH", "Golden Stores vs promedio = 100"),
                                    (fmt_c(sens[10]), "cajas / mes", "si P1 crece 10% (sensibilidad)")]):
        K.cifra(s, K.MARGEN + i * 2.3, y + 0.05, 2.15, v, e, dd, h=1.1, oscura=(i == 0))
    K.tarjeta(s, K.MARGEN, y + 1.35, 9.1, 1.35, "Decisión que pedimos hoy",
              f"1. Piloto de 90 días en {rec_corto} ({n_pil:,} tiendas) sorteado por zonas: mitad Atacar, mitad control (detecta desde {mde0:.0%}).\n"
              f"2. Que Bepensa entregue la venta mensual por tienda y las equivalencias de las cuentas sin ubicar ({sin_ubicar:.0f}% de las cajas).\n"
              f"3. Reactivar las {hl_na:,} tiendas HL que dejaron de comprar o están en riesgo antes de aplicar Bloquear.", acento=True)
    K.mensaje(s, "Fundamento: venta real de Bepensa (23 meses) × hogares 2025 a 300 m de cada tienda, con la metodología Golden Stores.")

    # 2 · por qué ahora: la mitad de las HL dejó de comprar (la brecha HL vs HH sale de la definición del clúster: no se usa)
    s, y = K.lamina(prs, f"{hl_na:,} de las {hl_n:,} tiendas HL tienen demanda alta alrededor pero dejaron de comprar o están en riesgo.",
                    "Por qué ahora",
                    notas=(f"Las HL tienen tantos hogares alrededor como las HH (mediana {hl_.hogares:,.0f} vs {hh_.hogares:,.0f}), pero solo el "
                           f"{act.loc['HL', 'activo']:.0f}% está activa (HH: {act.loc['HH', 'activo']:.0f}%). Ojo: la diferencia de venta HL vs HH "
                           f"({qa['brecha_real']:.0%}) no es una oportunidad medida: con venta al azar sale {qa['brecha_azar']:.0%}, porque el clúster se "
                           "define cortando por venta. Lo accionable es la actividad: una tienda que no compra no puede defender el anaquel."))
    estados = [e for e in ["activo", "en riesgo (1-2 m)", "inactivo (3+ m)"] if e in act.columns]
    K.barras(s, K.MARGEN, y, 5.4, 3.0, [f"{q} · {ACCION[q]}" for q in CLUSTERS], {e.capitalize(): act[e].values for e in estados},
             formato="0", apiladas=True, titulo="% de tiendas por actividad al cierre (ago-2026)", colores=[K.NEGRO, K.LIMA, K.GRIS_MEDIO])
    for i, (t, c) in enumerate([("Demanda alta, tienda apagada", f"HL: {act.loc['HL', 'activo']:.0f}% activas (HH {act.loc['HH', 'activo']:.0f}%) con "
                                                                 f"mediana de {hl_.hogares:,.0f} hogares a 300 m (HH {hh_.hogares:,.0f})."),
                                ("Aun activas, venden menos", f"Mediana de las activas: HL {qa['med_act']['HL']:.1f} cajas/mes vs HH {qa['med_act']['HH']:.1f}; "
                                                              f"competencia parecida ({comp['HL']:.0f} vs {comp['HH']:.0f} tiendas DENUE)."),
                                ("Venta concentrada", f"HH + LH = {hh_['% tiendas'] + lh_['% tiendas']:.0f}% de las tiendas y "
                                                      f"{hh_['% mix ventas'] + lh_['% mix ventas']:.0f}% de la venta.")]):
        K.tarjeta(s, K.MARGEN + 5.6, y + i * 1.0, 3.5, 0.9, t, c, oscura=(i == 0), compacta=True)
    K.mensaje(s, "Bloquear exige que la tienda compre: en HL, primero reactivar y después defender el anaquel.")

    K.notas(K.separador(prs, "La solución", "Golden Stores"), "Metodología NielsenIQ | Spectra aplicada tal cual: dos variables, cuatro clústeres, accionables y semáforos.")

    s, y = K.lamina(prs, "Cada tienda se mide con dos números: lo que vende y la demanda que tiene a 300 m.", "Metodología",
                    notas="Variable 1 viene de las ventas de Bepensa validadas en el notebook 00. Variable 2 es la demanda potencial del área transaccional, construida en el 04: Censo 2020 por manzana y localidades rurales, microsimulación, reparto por área en hexágonos H3 y actualización a 2025 con la Encuesta Intercensal 2025 (INEGI, sep-2026).")
    K.tarjeta(s, K.MARGEN, y + 0.1, 4.45, 2.5, "1 · Ventas de la tienda",
              "Cajas de sueros por mes de vida (oct-2024 → ago-2026), ventas Bepensa validadas en el notebook 00.\n"
              "Índice de ventas = venta de la tienda / promedio de su subcanal × 100.", oscura=True)
    K.tarjeta(s, K.MARGEN + 4.6, y + 0.1, 4.45, 2.5, "2 · Demanda potencial",
              f"Con base en su área transaccional: hogares 2025 a ≤ {C.RADIO_PDV_M} m (manzanas urbanas y localidades rurales repartidas por área) "
              "con su perfil por tamaño, niños, edad del jefe y NSE AMAI 2024, más competidores DENUE.\n"
              "Índice de demanda = hogares del área / promedio de su subcanal × 100.", acento=True)
    K.mensaje(s, f"Ejemplo: {e0.Nombre.title()} vende {e0['Cajas/mes (vida)']:.1f} cajas/mes (índice {e0.indice_venta:.0f}) y tiene "
                 f"{e0['HHs en el área']:,.0f} hogares a 300 m (índice {e0.indice_demanda:.0f}).")

    s, y = K.lamina(prs, "Venta × demanda separan las tiendas en cuatro clústeres con una acción distinta.", "Metodología",
                    notas="Alta = índice ≥ 100, es decir, sobre el promedio de su subcanal. Las Golden Stores son el clúster HH: alta demanda y alta venta.")
    x0, y0, wq, hq = K.MARGEN + 0.6, y + 0.05, 4.1, 1.25
    for q, (cx, cy) in {"HL": (0, 0), "HH": (1, 0), "LL": (0, 1), "LH": (1, 1)}.items():
        e = ej[q]
        K.tarjeta(s, x0 + cx * (wq + 0.1), y0 + cy * (hq + 0.1), wq, hq, f"{q} · {NOMBRE[q]} · {ACCION[q]}",
                  ("GOLDEN STORES\n" if q == "HH" else "") + f"Ej.: {e.Nombre.title()[:30]} · {e['Cajas/mes (vida)']:.1f} cajas/mes · "
                  f"{e['HHs en el área']:,.0f} hogares", oscura=(q == "HH"), compacta=True)
    for fila, et in enumerate(["Demanda\nalta", "Demanda\nbaja"]):
        K.texto(s, K.MARGEN, y0 + fila * (hq + 0.1) + 0.45, 0.55, 0.4, et, tam=7.5, negrita=True, color=K.GRIS)
    for colm, et in enumerate(["Ventas bajas", "Ventas altas"]):
        K.texto(s, x0 + colm * (wq + 0.1), y0 + 2 * hq + 0.15, wq, 0.18, et, tam=7.5, negrita=True, color=K.GRIS, alinear=K.PP_ALIGN.CENTER)
    K.mensaje(s, "Alta = índice ≥ 100 (sobre el promedio de su subcanal) · Golden Stores = HH.", invertido=True)

    s, y = K.lamina(prs, "Cada clúster tiene su acción: atacar, bloquear, fortalecer o mantener.", "Metodología",
                    notas="Accionables de la metodología Golden Stores, con una tienda real de cada clúster como ejemplo.")
    for q, (cx, cy) in {"HL": (0, 0), "HH": (1, 0), "LL": (0, 1), "LH": (1, 1)}.items():
        e = ej[q]
        K.tarjeta(s, K.MARGEN + cx * 4.6, y + 0.02 + cy * 1.6, 4.45, 1.48, f"{NOMBRE[q]} · {ACCION[q]}:",
                  "   ·   ".join(ACCIONABLES[q]) + f"\nEjemplo: {e.Nombre.title()[:34]} ({e.Subcanal.title()}, {e.Municipio})",
                  oscura=(q == "HH"), compacta=True)

    # ejemplo paso a paso
    s, y = K.lamina(prs, f"Ejemplo: {e0.Nombre.title()[:28]} es Golden Store en verde y su acción es atacar.", "Cómo se clasifica una tienda",
                    notas="Recorrido completo de una tienda real para que el equipo comercial entienda la lógica sin tecnicismos.")
    pasos = [("Venta", f"{e0['Cajas/mes (vida)']:.1f} cajas/mes · índice {e0.indice_venta:.0f} vs su subcanal ({e0.Subcanal.title()})"),
             ("Demanda", f"{e0['HHs en el área']:,.0f} hogares a 300 m · índice {e0.indice_demanda:.0f}"),
             ("Clúster", "Venta ≥ 100 y demanda ≥ 100 → HH (Golden Store)"),
             ("Semáforo", f"Sobre el promedio de su clúster en venta y demanda → {e0.semaforo}"),
             ("Acción", "Atacar: más SOS, exhibiciones, special packs, promoción"),
             ("Clase CP", "sin clase" if pd.isna(e0.clase_cp) else f"{e0.clase_cp} (potencial futuro)")]
    for i, (t, c) in enumerate(pasos):
        K.tarjeta(s, K.MARGEN + (i % 3) * 3.05, y + 0.05 + (i // 3) * 1.35, 2.9, 1.2, t, c, numero=i + 1, acento=(i == 4), compacta=True)
    K.mensaje(s, f"La misma lógica corre para las {n:,} tiendas: el Excel trae cada paso por tienda.")

    K.notas(K.separador(prs, "Evidencia", f"Canal {canal}"), "Resultados de la metodología sobre las tiendas tradicionales de la ZM.")

    s, y = K.lamina(prs, f"El {hh_['% tiendas']:.0f}% de las tiendas HH hacen el {hh_['% mix ventas']:.0f}% de las ventas de sueros en {canal}.",
                    "Panorama general",
                    notas=f"Lámina central de Golden Stores. HH: {hh_.tiendas:,.0f} tiendas, {hh_['% mix ventas']:.0f}% de la venta, índice {hh_['Index ventas']:.0f}. "
                          f"LH es el segundo clúster en venta ({lh_['% mix ventas']:.0f}%).")
    K.caja(s, K.MARGEN + 2.5, y - 0.05, 4.1, 0.3, relleno=K.NEGRO, redondeo=0.2)
    K.texto(s, K.MARGEN + 2.5, y + 0.02, 4.1, 0.18, f"Target: sueros · {n:,} tiendas", tam=9, negrita=True, color=K.LIMA, alinear=K.PP_ALIGN.CENTER)
    for q, (cx, cy) in {"HL": (0, 0), "HH": (1, 0), "LL": (0, 1), "LH": (1, 1)}.items():
        r = pan.loc[q]
        xx, yy = K.MARGEN + 0.35 + cx * 4.45, y + 0.4 + cy * 1.45
        K.caja(s, xx, yy, 0.5, 0.28, relleno=COLOR[q], redondeo=0.2)
        K.texto(s, xx, yy + 0.05, 0.5, 0.18, q, tam=10, negrita=True, color=K.LIMA if q == "HH" else K.BLANCO, alinear=K.PP_ALIGN.CENTER)
        K.texto(s, xx + 0.6, yy + 0.06, 3.5, 0.18, f"{NOMBRE[q]} · {ACCION[q]}", tam=9, negrita=True)
        for k, (et, val) in enumerate([("# Tiendas", f"{r.tiendas:,.0f} ({r['% tiendas']:.0f}%)"), ("% Mix ventas", f"{r['% mix ventas']:.0f}%"),
                                        ("Index ventas", f"{r['Index ventas']:.0f}")]):
            K.texto(s, xx + 0.3 + k * 1.3, yy + 0.42, 1.25, 0.15, et, tam=7, color=K.GRIS, alinear=K.PP_ALIGN.CENTER)
            K.texto(s, xx + 0.3 + k * 1.3, yy + 0.62, 1.25, 0.3, val, tam=15 if k == 2 else 13, negrita=(k == 2), alinear=K.PP_ALIGN.CENTER,
                    fuente=K.FUENTE_TITULO)
    for (xa, ya, xb, yb) in [(K.MARGEN + 4.55, y + 0.4, K.MARGEN + 4.55, y + 3.2), (K.MARGEN + 0.2, y + 1.78, K.MARGEN + 8.9, y + 1.78)]:
        ln = s.shapes.add_connector(1, K.Inches(xa), K.Inches(ya), K.Inches(xb), K.Inches(yb))
        ln.line.color.rgb = K.rgb(K.GRIS_CLARO)
    K.texto(s, K.MARGEN, y + 3.25, 9.1, 0.18, "↑ Demanda potencial (hogares a 300 m)          → Ventas (cajas / mes)", tam=7, color=K.GRIS)

    top_edad = pct.loc["HH", C.GRUPOS["Edad jefe"]].idxmax()
    ninos = pct.loc["HH", C.GRUPOS["Niños"]].idxmax()
    s, y = K.lamina(prs, f"Hogares de tiendas HH: índice A/B·C+ de {ab['HH']:.0f}, jefes de {top_edad.replace('Edad ', '').lower()} años y "
                         f"mayoría {ninos.lower()}.", "Perfil por clúster",
                    notas="Perfil de los hogares a 300 m de las tiendas de cada clúster. Resaltados los índices ≥ 105 vs el total de la ZM.")
    cats = [("Tamaño hogar", C.GRUPOS["Tamaño hogar"]), ("Presencia niños", C.GRUPOS["Niños"]), ("Edad jefe", C.GRUPOS["Edad jefe"]),
            ("NSE", C.GRUPOS["NSE AMAI 2024"])]
    xs = [K.MARGEN, K.MARGEN + 3.0, K.MARGEN + 5.05, K.MARGEN + 7.1]
    for i, q in enumerate(CLUSTERS):
        perfil_cluster(s, xs[i], y - 0.05, 3.0 if i == 0 else 2.0, q, b, pct, indice, cats, mostrar_etiquetas=(i == 0))
    K.texto(s, K.MARGEN, K.ALTO - 0.55, 9, 0.15, "% de hogares a ≤ 300 m de las tiendas del clúster · índice vs total de la ZM (resaltado ≥ 105).",
            tam=6.5, color=K.GRIS)

    s, y = K.lamina(prs, f"Las Golden Stores se concentran en {mun_top}: {mm.loc[mun_top, 'HH']:.0f}% de las tiendas HH.", "Perfil por clúster",
                    notas="Distribución de las tiendas de cada clúster por municipio y subcanal.")
    tb = mm.round(0).reset_index().rename(columns={"Municipio": "Municipio (% de tiendas)"})
    K.tabla(s, tb, K.MARGEN, y + 0.05, 4.4, anchos=[1.6, 0.7, 0.7, 0.7, 0.7], alto_fila=0.3, formato={q: (lambda v: f"{v:.0f}%") for q in CLUSTERS})
    ms = (pd.crosstab(b.Subcanal.str.title(), b.cluster, normalize="columns") * 100).reindex(columns=CLUSTERS).fillna(0)
    tb2 = ms.round(0).reset_index().rename(columns={"Subcanal": "Subcanal (% de tiendas)"})
    K.tabla(s, tb2, K.MARGEN + 4.7, y + 0.05, 4.4, anchos=[1.6, 0.7, 0.7, 0.7, 0.7], alto_fila=0.3, formato={q: (lambda v: f"{v:.0f}%") for q in CLUSTERS})

    s, y = K.lamina(prs, f"El CP ordena el potencial futuro: {cpx.loc['HH', ['Very High', 'High']].sum():.0f}% de las HH son High o Very High.",
                    "Clasificación CP", "Clase de potencial del Customer Potential (Very High → Low). Clasifica dentro del clúster; no decide sola.",
                    notas="El CP mide cuánto más puede crecer cada cliente frente a su comparable. Lo usamos para ordenar dentro de cada clúster, no para decidir el clúster.")
    K.barras(s, K.MARGEN, y, 5.8, 3.0, [f"{q} · {ACCION[q]}" for q in CLUSTERS], {c: cpx[c].values for c in ORDEN_CP}, formato="0;;;",
             titulo="% de tiendas por clase de potencial CP", apiladas=True, colores=[K.LIMA, K.VERDE, K.GRIS_MEDIO, K.GRIS_CLARO])
    K.tarjeta(s, K.MARGEN + 6.0, y + 0.1, 3.1, 2.8, "Lectura",
              f"High + Very High: HH {cpx.loc['HH', ['Very High', 'High']].sum():.0f}% · LH {cpx.loc['LH', ['Very High', 'High']].sum():.0f}% · "
              f"HL {cpx.loc['HL', ['Very High', 'High']].sum():.0f}% · LL {cpx.loc['LL', ['Very High', 'High']].sum():.0f}%.\n"
              "Dentro de cada clúster, empezar por las tiendas con clase CP más alta.", oscura=True)

    verdes = int(sem["verde"].sum())
    s, y = K.lamina(prs, f"Los semáforos dejan {verdes:,} tiendas en verde para actuar primero.", "Priorización",
                    "Semáforo individualizado por clúster (HH, HL, LH, LL) con la desviación estándar de venta y demanda dentro del clúster.",
                    notas="Verde: venta y demanda sobre el promedio de su clúster. Rojo: ambas debajo. Amarillo: alrededor del promedio o con una de las dos subdesarrollada.")
    for i, (col, t, c) in enumerate([("verde", "Verde", "Tiendas con altas ventas y alto índice de demanda: se gasta más y tienen mayor potencial."),
                                     ("amarillo", "Amarillo", "Alrededor de los promedios (#2) + subdesarrollo en ventas o demanda: #1 alta venta / demanda baja · "
                                                              "#3 demanda alta / venta menor."),
                                     ("rojo", "Rojo", "Tiendas con bajo índice de demanda y bajas ventas promedio: se vende menos y el potencial es bajo.")]):
        yy = y + 0.1 + i * 0.95
        K.caja(s, K.MARGEN, yy + 0.1, 0.5, 0.5, relleno=SEMAFORO[col], forma=K.MSO_SHAPE.OVAL)
        K.texto(s, K.MARGEN + 0.7, yy + 0.08, 8.3, 0.22, t, tam=11, negrita=True, fuente=K.FUENTE_TITULO)
        K.texto(s, K.MARGEN + 0.7, yy + 0.34, 8.3, 0.5, c, tam=9, color=K.GRIS)
    K.mensaje(s, "Verde: z de venta y de demanda (ambas en log) > 0 · Rojo: ambos < 0 · Amarillo: |z| ≤ 0.5 en ambos o solo uno arriba.")

    s, y = K.lamina(prs, f"{p1_n:,} tiendas P1 (HH y LH en verde y amarillo) son la prioridad de la categoría.", "Priorización",
                    f"Target: sueros · {n:,} tiendas. P1: capitalizar con los clústeres de alta venta de la categoría.",
                    notas=f"P1 = {p1_n:,} tiendas: {p1_mix:.0f}% de la venta. P1 conserva el {qa['p1_de_hhlh']:.0%} de las HH + LH: el semáforo filtra poco, "
                          "por eso dentro de P1 conviene empezar por verde.")
    tb = sem.reset_index().rename(columns={"cluster": "Clúster", "verde": "Verde", "amarillo": "Amarillo", "rojo": "Rojo"})
    tb["Clúster"] = tb["Clúster"] + " · " + tb["Clúster"].map(ACCION)
    K.tabla(s, tb, K.MARGEN, y + 0.1, 5.2, anchos=[1.9, 1.1, 1.1, 1.1], alto_fila=0.42, tam=10, resaltar=[0, 2])
    K.cifra(s, K.MARGEN + 5.6, y + 0.1, 3.5, f"{p1_n:,}", "tiendas P1", f"{p1_mix:.0f}% de la venta del canal", h=1.2, oscura=True)
    K.tarjeta(s, K.MARGEN + 5.6, y + 1.45, 3.5, 1.3, "Siguiente paso", f"Lista en el Excel (hoja Tiendas, filtro P1). {p1_front:,} tiendas P1 están a "
              "±10 puntos del umbral (columna Frontera): revisarlas con el equipo antes de rutear.", compacta=True)
    K.mensaje(s, "P1 · Capitalizar con los clústeres de alta venta: HH y LH en semáforo verde y amarillo.", invertido=True)

    top = (b[(b.cluster == "HH") & (b.semaforo == "verde")].sort_values("Cajas/mes (vida)", ascending=False)
           .head(10)[["Nombre", "Subcanal", "Municipio", "Cajas/mes (vida)", "HHs en el área", "indice_demanda", "clase_cp"]])
    s, y = K.lamina(prs, "Estas 10 Golden Stores en verde son las primeras a atacar.", "Priorización",
                    "Tiendas HH con semáforo verde, ordenadas por venta. La clase CP acompaña la clasificación.",
                    notas="Ejemplos concretos para el equipo de campo: las tiendas HH verde de mayor venta.")
    tt = top.rename(columns={"Cajas/mes (vida)": "Cajas/mes", "HHs en el área": "Hogares 300 m", "indice_demanda": "Índice demanda", "clase_cp": "Clase CP"})
    tt["Nombre"] = tt.Nombre.str.slice(0, 32)
    tt["Subcanal"] = tt.Subcanal.str.title()
    K.tabla(s, tt, K.MARGEN, y, 9.1, anchos=[2.6, 1.55, 1.05, 0.85, 1.05, 1.0, 1.0], alto_fila=0.26, tam=7, resaltar=[0, 1, 2],
            formato={"Cajas/mes": lambda v: f"{v:,.1f}", "Hogares 300 m": lambda v: f"{v:,.0f}", "Índice demanda": lambda v: f"{v:.0f}",
                     "Clase CP": lambda v: "sin clase" if pd.isna(v) else str(v)})
    K.mensaje(s, "Atacar: más SOS, exhibiciones, special packs, ítems exclusivos, precio, comunicación y promoción.")

    # caso de negocio (en cajas: no hay precio ni costo en los datos)
    s, y = K.lamina(prs, f"Cada +10% de venta en las tiendas P1 suma {fmt_c(sens[10])} cajas/mes al canal ({sens[10] / total_cajas:+.1%}).",
                    "Caso de negocio", "Sensibilidad (no pronóstico) sobre la venta actual de las tiendas P1. Sin precio ni costo en los datos: se expresa en cajas.",
                    notas="No proyectamos pesos porque no tenemos precio ni costo de servir. La sensibilidad muestra el tamaño del premio si el piloto funciona.")
    K.barras(s, K.MARGEN, y, 5.2, 2.9, [f"+{p}% en P1" for p in sens], {"cajas / mes adicionales": list(sens.values())}, formato="#,##0",
             titulo="Cajas / mes adicionales", horizontal=False, colores=[K.NEGRO], leyenda=False)
    for i, (v, e, dd) in enumerate([(fmt_c(p1_cajas), "cajas / mes en P1 hoy", f"{p1_mix:.0f}% del canal"),
                                    (fmt_c(total_cajas), "cajas / mes del canal", f"{n:,} tiendas analizadas"),
                                    (f"{n_pil:,}", "tiendas en el piloto", f"{rec_corto} · sorteo por zonas")]):
        K.cifra(s, K.MARGEN + 5.5, y + i * 0.99, 3.6, v, e, dd, h=0.94, oscura=(i == 0))
    K.mensaje(s, "Se mide piloto vs control sorteado, nunca antes vs después: las tiendas elegidas por venta alta bajan solas.")

    # ruta de ejecución y puertas de decisión
    s, y = K.lamina(prs, f"Un piloto de 90 días sorteado por zonas decide si se escala a todas las P1.", "Ruta de ejecución",
                    notas=(f"Piloto en {rec} ({n_pil:,} tiendas, {int(fr['zonas H3 r7'])} zonas H3 de ≈ 5 km²). Efecto mínimo detectable {mde0:.0%} sin línea "
                           f"base; con venta mensual (línea base) bajaría a ≈ {mde8:.0%}. Por eso pedimos la venta mensual en el mes 1. "
                           f"El {fr['control a < 600 m de una tratada']:.0%} de los controles queda a < 600 m de una tratada: en la medición se "
                           "excluyen esos controles (sus áreas de 300 m se mezclan)."))
    fases = [("Mes 1 · Preparar y lanzar", "Recibir venta mensual por tienda y equivalencias de cuentas.\nSortear zonas H3: mitad Atacar, mitad control "
                                          f"(hoja Piloto del Excel). Reactivar las {hl_na:,} HL que no compran."),
             ("Mes 2 · Medir", f"Nuevo corte de ventas → correr los notebooks 00 → 05.\nPuerta 1: ¿piloto supera al control por ≥ {mde0:.0%}?"),
             ("Mes 3 · Escalar", f"Si pasa la puerta 1: extender a las {p1_n:,} P1 y activar Bloquear en las HL activas.\nPuerta 2: incremento vs costo de servir.")]
    for i, (t, c) in enumerate(fases):
        K.tarjeta(s, K.MARGEN + i * 3.05, y + 0.1, 2.9, 2.5, t, c, numero=i + 1, oscura=(i == 1), compacta=True)
    K.mensaje(s, "Cada puerta permite detener o ajustar sin comprometer toda la red de tiendas.", invertido=True)

    # riesgos y mitigación
    riesgos = [("Llave ventas ↔ CP inferida", "El id de ventas se cruzó por prefijo + id del CP (73% de empate, probado estadísticamente).",
                "Bepensa confirma la tabla de equivalencias en el mes 1."),
               ("Volumen sin ubicar", f"{sin_ubicar:.0f}% de las cajas (pocas cuentas grandes) no tiene coordenadas en el CP.",
                "Pedir coordenadas de esas cuentas y reclasificar."),
               ("Tiendas sin hogares a 300 m", f"{qa['excl_n']:,} tiendas ({qa['excl_mix']:.0%} de la venta del canal) en zonas comerciales quedan fuera de los clústeres.",
                "Van en hoja aparte del Excel para tratarlas por venta y competencia."),
               ("La demanda no pronostica", f"Hogares a 300 m vs venta de la tienda: ρ = {qa['rho_demanda']:+.2f}. La demanda clasifica el área, no predice la venta.",
                "Validar con el piloto; sumar tráfico o puntos de interés si se consiguen."),
               ("Hogares 2020 → 2025", "Estructura por manzana del Censo 2020, actualizada a 2025 por municipio con la Intercensal 2025.",
                "Las zonas nuevas dentro de cada municipio siguen subrepresentadas."),
               ("Clase CP y tamaño", "La clase CP mide potencial absoluto (venta × brecha): las clases altas solo aparecen en tiendas que ya venden.",
                "Leerla junto a la brecha relativa (columna del Excel).")]
    s, y = K.lamina(prs, f"Los {len(riesgos)} riesgos del análisis tienen mitigación concreta antes del piloto.", "Riesgos",
                    notas="Reconocer los límites da credibilidad. Todos se recalculan en el notebook 05 (sección 2b).")
    tbr = pd.DataFrame(riesgos, columns=["Riesgo", "Qué pasa", "Mitigación"])
    K.tabla(s, tbr, K.MARGEN, y + 0.05, 9.1, anchos=[2.0, 4.1, 3.0], alto_fila=0.44, tam=7.5)
    K.mensaje(s, "Se clasifica y se valida con piloto; no se pronostica la venta de cada tienda.")

    s, y = K.lamina(prs, "Cada tienda sale con su clúster, semáforo y acción en un Excel listo para ruteo.", "Entregable",
                    notas="El Excel corporativo trae portada, resumen con fórmulas, tiendas con filtros y colores, perfiles, accionables, guion y diccionario.")
    for i, (t, c) in enumerate([(f"05_golden_stores_{canal.lower()}_*.xlsx", "Resumen con fórmulas · Tiendas con clúster, semáforo, P1 y coordenadas · Perfiles · Accionables · Guion."),
                                (f"04_pdv_*_{canal.lower()}_hexagonos_*.xlsx", "Perfil completo por tienda: HHs · % · Índice de todas las categorías y competencia a 300 m."),
                                ("00_validacion_ventas_cp_*.xlsx", "Calidad de los datos: llave ventas ↔ CP, señales útiles y descartadas, sesgos."),
                                ("hexagonos_h3r10_*.gpkg", "Hexágonos con hogares, NSE y red Bepensa para mapas.")]):
        K.tarjeta(s, K.MARGEN, y + 0.05 + i * 0.8, 9.1, 0.72, t, c, oscura=(i == 0), compacta=True)

    # anexo
    K.notas(K.separador(prs, "Anexo", "Metodología y datos"), "Láminas de respaldo para preguntas.")
    s, y = K.lamina(prs, "El modelo sigue seis pasos: del target a geolocalizar cada tienda.", "Anexo · metodología",
                    notas="Pasos de la metodología Golden Stores (NielsenIQ | Spectra).")
    for i, (t, c) in enumerate([("Target", f"Sueros en el canal {canal}"), ("Competidores", "Densidad comercial DENUE a 300 m"),
                                ("Buscar ganancias", "Tiendas de alta venta y alta demanda"), ("Optimizar recursos", "Acción distinta por clúster"),
                                ("Identificar prioridades", "Semáforos dentro de cada clúster"), ("Geolocalizar tiendas", "Coordenadas y hexágonos H3 por tienda")]):
        K.tarjeta(s, K.MARGEN + (i % 3) * 3.05, y + 0.05 + (i // 3) * 1.25, 2.9, 1.1, t, c, numero=i + 1, compacta=True)
    s, y = K.lamina(prs, f"El análisis cubre {n:,} tiendas tradicionales con 23 meses de venta real.", "Anexo · set up",
                    notas="Alcance, periodo y limitaciones declaradas.")
    setup = [("Mercado", f"Canal {canal} · {C.ZM_NOMBRE} ({', '.join(C.ZM_MUNICIPIOS.values())})"), ("Periodo", "Oct-2024 → ago-2026 (23 meses)"),
             ("Target", f"Sueros · {n:,} tiendas con venta y hogares 2025 a ≤ {C.RADIO_PDV_M} m (urbanos y rurales)"),
             ("Líneas de reporte", "Total sueros (sin SKU ni marca: share, distribución y precio no aplican)"),
             ("Semáforos", "Individualizados por clúster, con desviación estándar de venta y demanda"),
             ("Clasificación CP", "Clase de potencial futuro (Very High → Low) como variable adicional")]
    for i, (t, c) in enumerate(setup):
        K.tarjeta(s, K.MARGEN + (i % 2) * 4.6, y + 0.02 + (i // 2) * 1.0, 4.45, 0.88, t, c, oscura=(i == 2), compacta=True)
    s, y = K.lamina(prs, "Las señales del cliente se depuraron: dos columnas del CP eran la propia venta.", "Anexo · notebook 00",
                    notas="Validación de datos: llave, fuga de información y actividad.")
    for i, (t, c) in enumerate([("Llave inferida", h("Llave").split(";")[0] + "."),
                                ("Fuga: no usar", "PotentialQuantitative es la venta actual (99% idéntica) y size_class se reconstruye con la venta (97.5%)."),
                                ("CP: potencial futuro", "PotentialQuantitativeFinal y su clase miden potencial absoluto (venta × brecha): clasifican. La brecha relativa compara tiendas de distinto tamaño."),
                                ("Actividad", h("Actividad").split("Kaplan")[0].strip())]):
        K.tarjeta(s, K.MARGEN + (i % 2) * 4.6, y + 0.05 + (i // 2) * 1.35, 4.45, 1.22, t, c, oscura=(i == 1))
    rr = rob[rob.robusta & rob.subcanal.str.startswith(canal)]
    hh_neg = rr[rr.indicador == "HHs en el área"]
    s, y = K.lamina(prs, "El entorno explica poco de la venta individual: por eso se clasifica, no se pronostica.", "Anexo · notebook 04",
                    "Spearman dentro de cada segmento; robusta = sobrevive a bootstrap espacial por bloques, solo activos y sin atípicos.",
                    notas="Respaldo estadístico para la pregunta '¿qué tanto explica el entorno la venta?'.")
    rs = rr.reindex(rr.spearman.abs().sort_values(ascending=False).index).head(8)
    et = [f"{'(+)' if v > 0 else '(−)'} {i_} · {sc.split(' · ')[-1].title()}" for i_, sc, v in zip(rs.indicador, rs.subcanal, rs.spearman)]
    K.barras(s, K.MARGEN, y, 6.2, 3.0, et, {"positiva": np.where(rs.spearman > 0, rs.spearman.abs(), 0),
                                            "negativa": np.where(rs.spearman < 0, rs.spearman.abs(), 0)},
             formato="0.00;;;", apiladas=True, colores=[K.LIMA, K.GRIS_MEDIO], leyenda=True)
    K.tarjeta(s, K.MARGEN + 6.4, y + 0.1, 2.7, 2.7, "Lectura", f"{len(rr)} asociaciones robustas en {canal}, todas con |ρ| ≤ {rr.spearman.abs().max():.2f}. "
              + (f"Más hogares a 300 m = menos venta por tienda en {int((hh_neg.spearman < 0).sum())} de {len(hh_neg)} subcanales (competencia)." if len(hh_neg) else ""))

    objeciones = [("¿Por qué no priorizar solo con el CP?", "Porque el CP mide potencial futuro frente a un comparable; la venta real y la demanda "
                                                           "alrededor definen el clúster. El CP ordena dentro de cada clúster."),
                  ("¿Qué tan segura es la llave entre ventas y CP?", f"Empata el 73% con prueba estadística (azar = 4%); se confirma con Bepensa en el mes 1 "
                                                                     "antes de escalar."),
                  (f"¿Y el {sin_ubicar:.0f}% de volumen que no aparece?", "Son pocas cuentas grandes sin coordenadas; no cambian la clasificación de las tiendas "
                                                                          "analizadas y se integran cuando llegue su tabla."),
                  ("¿Por qué las HL venden menos si tienen la misma demanda?", f"{hl_na / hl_n:.0%} dejó de comprar o está en riesgo. La brecha HL–HH sola no "
                                                                              "mide oportunidad: sale de cómo se define el clúster.")]
    s, y = K.lamina(prs, f"Las {len(objeciones)} preguntas difíciles ya tienen respuesta.", "Anexo · objeciones (pre-mortem)",
                    notas="Técnica Asegurar + Responder + Anclar: responder con sí/no y una cifra, y volver al flujo.")
    for i, (qq, rsp) in enumerate(objeciones):
        K.tarjeta(s, K.MARGEN, y + 0.02 + i * 0.86, 9.1, 0.78, qq, rsp, oscura=(i == 0), compacta=True)

    K.notas(K.cierre(prs, "Let knowledge in.", f"Kin Analytics · Golden Stores · Sueros · Canal {canal} · Bepensa"), "Cierre.")
    guion = [("Elevator hook (60 s)", [prs.slides[1].notes_slide.notes_text_frame.text.replace("GUION 60 s: ", "")]),
             ("Objeciones críticas (pre-mortem)", [f"{qq} → {rsp}" for qq, rsp in objeciones]),
             ("Riesgos y mitigación", [f"{r[0]}: {r[1]} Mitigación: {r[2]}" for r in riesgos]),
             ("Ruta de ejecución", [f"{t}: {c}".replace("\n", " ") for t, c in fases])]
    return prs, p1, guion


def excel_corporativo(canal, b, pan, pct, indice, sem, p1, guion, ruta):
    """Entregable Excel con formato Kin: portada, resumen con fórmulas, tiendas, perfiles, accionables, guion y diccionario."""
    import openpyxl
    import excel_kin as X
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pil = PILOTO[canal]["asignacion"].set_index("pos_id_cp")
    tiendas = (b.assign(P1=b.pos_id_cp.isin(p1.pos_id_cp).map({True: "Sí", False: "No"}),
                        Frontera=b.frontera.map({True: "Sí", False: "No"}), **{"Alta reciente": b.alta_reciente.map({True: "Sí", False: "No"})},
                        **{"Grupo piloto": b.pos_id_cp.map(pil.grupo).fillna("—"), "Zona piloto (H3 r7)": b.pos_id_cp.map(pil.zona).fillna("—")})
               .rename(columns={"pos_id_cp": "ID PDV", "cluster": "Clúster", "accion": "Acción", "semaforo": "Semáforo", "semaforo_detalle": "Detalle semáforo",
                                "Cajas/mes (vida)": "Cajas/mes", "indice_venta": "Índice venta", "HHs en el área": "Hogares 300 m",
                                "indice_demanda": "Índice demanda", "clase_cp": "Clase CP", "Estado actividad": "Actividad",
                                "DENUE Moderno en el área": "DENUE moderno 300 m", "DENUE Tradicional en el área": "DENUE tradicional 300 m",
                                "brecha_relativa": "Brecha relativa CP"}))
    cols = ["ID PDV", "Nombre", "Subcanal", "Municipio", "Localidad", "Clúster", "Acción", "Semáforo", "Detalle semáforo", "P1", "Frontera", "Cajas/mes",
            "Índice venta", "Hogares 300 m", "Índice demanda", "Clase CP", "Brecha relativa CP", "Actividad", "Alta reciente", "Grupo piloto",
            "Zona piloto (H3 r7)", "DENUE moderno 300 m", "DENUE tradicional 300 m", "Latitud", "Longitud"]
    tiendas = tiendas[cols].sort_values(["Clúster", "Semáforo", "Cajas/mes"], ascending=[True, False, False])
    tiendas["Subcanal"] = tiendas.Subcanal.str.title()
    X.hoja_tabla(wb, "Tiendas", tiendas, formatos={"Cajas/mes": "#,##0.0", "Índice venta": "0", "Hogares 300 m": "#,##0", "Índice demanda": "0",
                                                  "Brecha relativa CP": "0.00",
                                                  "Latitud": "0.000000", "Longitud": "0.000000"},
                 semaforo_col="Semáforo", cluster_col="Clúster", barras=["Cajas/mes", "Hogares 300 m"],
                 anchos={"Nombre": 34, "Detalle semáforo": 26, "Acción": 11})
    # resumen con fórmulas vivas sobre la hoja Tiendas
    ws = wb.create_sheet("Resumen", 0)
    ws.sheet_view.showGridLines = False
    n_fin = len(tiendas) + 1
    col = {c: openpyxl.utils.get_column_letter(i + 1) for i, c in enumerate(cols)}
    rng = lambda c: f"Tiendas!${col[c]}$2:${col[c]}${n_fin}"
    ws["B2"] = f"Golden Stores · Sueros · Canal {canal} · {C.ZM_NOMBRE}"
    ws["B2"].font = Font(name=X.F, size=14, bold=True)
    ws["B3"] = "Fórmulas vivas sobre la hoja Tiendas: si se filtra o corrige una tienda, el resumen se actualiza."
    ws["B3"].font = Font(name=X.F, size=9, italic=True, color=X.GRIS)
    enc = ["Clúster", "Acción", "# Tiendas", "% Tiendas", "Cajas / mes", "% Mix ventas", "Index ventas", "Verde", "Amarillo", "Rojo", "Mediana hogares 300 m"]
    for j, e in enumerate(enc, 2):
        c = ws.cell(5, j, e)
        c.font = Font(name=X.F, size=9, bold=True, color=X.BLANCO)
        c.fill = X._relleno(X.NEGRO)
    for i, q in enumerate(CLUSTERS, 6):
        ws.cell(i, 2, q)
        ws.cell(i, 3, ACCION[q])
        ws.cell(i, 4, f'=COUNTIF({rng("Clúster")},B{i})')
        ws.cell(i, 5, f"=D{i}/SUM($D$6:$D$9)")
        ws.cell(i, 6, f'=SUMIF({rng("Clúster")},B{i},{rng("Cajas/mes")})')
        ws.cell(i, 7, f"=F{i}/SUM($F$6:$F$9)")
        ws.cell(i, 8, f"=(F{i}/D{i})/(SUM($F$6:$F$9)/SUM($D$6:$D$9))*100")
        for k, semv in enumerate(["verde", "amarillo", "rojo"]):
            ws.cell(i, 9 + k, f'=COUNTIFS({rng("Clúster")},B{i},{rng("Semáforo")},"{semv}")')
        ws.cell(i, 12, float(pan.loc[q, "hogares"]))
    ws.cell(10, 2, "Total")
    for j, letra in [(4, "D"), (6, "F"), (9, "I"), (10, "J"), (11, "K")]:
        ws.cell(10, j, f"=SUM({letra}6:{letra}9)")
    ws.cell(10, 5, "=SUM(E6:E9)")
    ws.cell(10, 7, "=SUM(G6:G9)")
    X.estilo_formulas(ws, "B6:L10")
    X.estilo_formulas(ws, "E6:E10", "0.0%")
    X.estilo_formulas(ws, "G6:G10", "0.0%")
    X.estilo_formulas(ws, "F6:F10", "#,##0")
    X.estilo_formulas(ws, "H6:H9", "0")
    X.estilo_formulas(ws, "L6:L9", "#,##0")
    X.estilo_formulas(ws, "B10:L10", negrita=True, relleno=X.FONDO)
    ws.cell(12, 2, "P1 · HH y LH en verde y amarillo").font = Font(name=X.F, size=10, bold=True)
    ws.cell(12, 4, f'=COUNTIF({rng("P1")},"Sí")')
    ws.cell(12, 6, f'=SUMIF({rng("P1")},"Sí",{rng("Cajas/mes")})')
    ws.cell(12, 7, "=F12/F10")
    X.estilo_formulas(ws, "B12:G12", negrita=True, relleno=X.LIMA)
    X.estilo_formulas(ws, "F12:F12", "#,##0", negrita=True, relleno=X.LIMA)
    X.estilo_formulas(ws, "G12:G12", "0.0%", negrita=True, relleno=X.LIMA)
    ws.cell(13, 2, "Mediana de hogares: valor calculado en el notebook 05 (no es fórmula).").font = Font(name=X.F, size=8, italic=True, color=X.GRIS)
    for j, w in zip(range(2, 13), [12, 12, 11, 11, 12, 12, 12, 9, 10, 9, 20]):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.column_dimensions["A"].width = 3

    X.hoja_tabla(wb, "Perfil %", pct.round(1).rename_axis("Clúster"), indice=True, titulo="Perfil de hogares a ≤ 300 m (% por grupo)",
                 nota="Cada grupo (tamaño, niños, edad del jefe, NSE) suma 100 dentro de cada clúster.", formatos={c: "0.0" for c in C.CATEGORIAS})
    X.hoja_tabla(wb, "Perfil índice", indice.round(0).rename_axis("Clúster"), indice=True, titulo="Índice vs total de la ZM (100 = promedio)",
                 nota="≥ 105 = sobrerrepresentado; ≤ 95 = subrepresentado.", formatos={c: "0" for c in C.CATEGORIAS})
    X.hoja_texto(wb, "Accionables", [(f"{q} · {NOMBRE[q]} · {ACCION[q]}", ACCIONABLES[q]) for q in CLUSTERS] +
                 [("Semáforos", ["Verde: venta y demanda sobre el promedio del clúster.", "Rojo: ambas bajo el promedio del clúster.",
                                 "Amarillo: alrededor de los promedios (#2) o subdesarrollo en una (#1 alta venta / demanda baja, #3 demanda alta / venta menor).",
                                 "P1: capitalizar con HH y LH en verde y amarillo."])])
    X.hoja_texto(wb, "Guion y objeciones", guion)
    tp = PILOTO[canal]["tabla"].copy()
    X.hoja_tabla(wb, "Piloto", tp, titulo=f"Diseño del piloto · recomendado: {PILOTO[canal]['recomendado']}",
                 nota="Sorteo por zonas H3 res 7 emparejadas por tamaño (semilla fija). MDE = efecto mínimo detectable (α 5%, potencia 80%) con efecto de "
                      "diseño por zona; 'con línea base' supone correlación 0.8 con la venta previa (requiere venta mensual). La columna Grupo piloto "
                      "de la hoja Tiendas trae la asignación.",
                 formatos={"ICC zona": "0.000", "efecto de diseño": "0.00", "MDE sin línea base": "0.0%", "MDE con línea base (ρ=0.8)": "0.0%",
                           "control a < 600 m de una tratada": "0.0%"}, anchos={"universo": 38})
    ex = EXCL[canal][["pos_id_cp", "Nombre", "Subcanal", "Municipio", "Localidad", "Cajas/mes (vida)", "Estado actividad", "Latitud", "Longitud"]].rename(
        columns={"pos_id_cp": "ID PDV", "Cajas/mes (vida)": "Cajas/mes", "Estado actividad": "Actividad"}).sort_values("Cajas/mes", ascending=False)
    X.hoja_tabla(wb, "Sin hogares 300 m", ex, titulo=f"Tiendas {canal} con venta y sin hogares a ≤ {C.RADIO_PDV_M} m (fuera de los clústeres)",
                 nota="Zonas comerciales o sin manzanas habitadas: la demanda por hogares no aplica. Tratarlas por venta y competencia.",
                 formatos={"Cajas/mes": "#,##0.0", "Latitud": "0.000000", "Longitud": "0.000000"}, anchos={"Nombre": 34})
    dic = pd.DataFrame([("Cajas/mes", "Cajas de sueros por mes de vida de la tienda (oct-2024 → ago-2026)."),
                        ("Índice venta", "Venta de la tienda / promedio de su subcanal × 100."),
                        ("Hogares 300 m", "Hogares 2025 a ≤ 300 m (hexágonos H3 res 10): manzanas urbanas (Censo 2020 + microsimulación) y localidades rurales "
                                          "(ITER 2020) repartidas por área, × crecimiento de población 2020→2025 del municipio (Intercensal 2025)."),
                        ("Frontera", "Sí = a ±10 puntos del índice 100 en venta o demanda: su clúster cambia con poco ruido."),
                        ("Brecha relativa CP", "PotentialEstimatedToCover: brecha vs el comparable, independiente del tamaño de la tienda."),
                        ("Alta reciente", "Sí = 3 meses de vida o menos: su venta mensual todavía es ruidosa."),
                        ("Grupo piloto / Zona piloto", "Asignación del sorteo por zonas H3 res 7 para el universo recomendado del piloto (hoja Piloto)."),
                        ("Índice demanda", "Hogares del área / promedio de su subcanal × 100."),
                        ("Clúster", "HH, HL, LH, LL: demanda (primera letra) × venta (segunda letra); alta = índice ≥ 100."),
                        ("Semáforo", "Individualizado por clúster con la desviación estándar de venta y demanda (ambas en log)."),
                        ("P1", "HH y LH en semáforo verde o amarillo."),
                        ("Clase CP", "Clase de potencial futuro del Customer Potential (Very High → Low). Clasifica, no decide."),
                        ("DENUE 300 m", "Tiendas DENUE 05_2026 (moderno / tradicional) en el área de la tienda.")], columns=["Campo", "Definición"])
    X.hoja_tabla(wb, "Diccionario", dic, anchos={"Campo": 22, "Definición": 100})
    X.portada(wb, "Golden Stores", f"Sueros · Canal {canal} · {C.ZM_NOMBRE} · Bepensa",
              [("Metodología", "Golden Stores (NielsenIQ | Spectra): ventas × demanda potencial → clústeres HH / HL / LH / LL, accionables y semáforos."),
               ("Universo", f"{len(b):,} tiendas del canal {canal} con venta y hogares a ≤ {C.RADIO_PDV_M} m."),
               ("Periodo", "Ventas oct-2024 → ago-2026 · Censo 2020 + Intercensal 2025 · Marco Geoestadístico Intercensal 2025 · DENUE 05_2026."),
               ("Elaboró", "Kin Analytics · generado por notebooks/_src/05_presentacion_pdv.py")],
              [("Resumen", "Clústeres, mix de ventas, índice y semáforos con fórmulas vivas."), ("Tiendas", "Una fila por tienda: clúster, acción, semáforo, P1, índices, clase CP, coordenadas."),
               ("Perfil %", "Perfil de hogares por clúster."), ("Perfil índice", "Índices vs total de la ZM."), ("Accionables", "Acciones por clúster y reglas de semáforo."),
               ("Guion y objeciones", "Elevator hook, objeciones, riesgos y ruta de ejecución."), ("Piloto", "Opciones de piloto, MDE y sorteo por zonas."),
               ("Sin hogares 300 m", "Tiendas con venta fuera de los clústeres (zonas comerciales)."), ("Diccionario", "Definición de cada columna.")])
    wb.save(ruta)


for canal, (b, pan, pct, indice, sem) in res.items():
    prs, p1, guion = deck(canal, b, pan, pct, indice, sem)
    out_ppt = C.OUT / f"05_golden_stores_{canal.lower()}_{C.CLIENTE}_{C.SLUG}.pptx"
    prs.save(out_ppt)
    out_x = C.OUT / f"05_golden_stores_{canal.lower()}_{C.CLIENTE}_{C.SLUG}.xlsx"
    excel_corporativo(canal, b, pan, pct, indice, sem, p1, guion, out_x)
    print(f"Guardado: {out_ppt.name} ({len(prs.slides)} láminas) | {out_x.name} ({len(b):,} tiendas, {len(p1):,} P1)")
    print("\n".join(f"· {t}: {' | '.join(p)[:300]}" for t, p in guion[:1]))
